"""Phase 0 inventory: parse EP1581 reference report → YAML spec."""

from __future__ import annotations

import re
from pathlib import Path

import yaml
from docx import Document

from ansys_report.ep1581_paths import ep1581_reference_docx

OUT = Path(__file__).resolve().parents[1] / "docs" / "ep1581_data_spec.yaml"


def parse_reference_docx(path: Path) -> dict:
    doc = Document(str(path))
    sections = [
        {"level": p.style.name, "title": p.text.strip()}
        for p in doc.paragraphs
        if p.style.name.startswith("Heading") and p.text.strip()
    ]
    captions = [p.text.strip() for p in doc.paragraphs if p.style.name == "Caption" and p.text.strip()]
    figures = [c for c in captions if c.lower().startswith("figure")]
    table_captions = [c for c in captions if c.lower().startswith("table")]

    tables = []
    for ti, table in enumerate(doc.tables):
        first_row = [c.text.strip()[:80] for c in table.rows[0].cells] if table.rows else []
        tables.append({"index": ti + 1, "rows": len(table.rows), "first_row": first_row})

    return {
        "path": str(path.name),
        "stats": {
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "figure_captions": len(figures),
            "table_captions": len(table_captions),
            "headings": len(sections),
        },
        "sections": sections,
        "figure_captions": figures[:50],
        "table_captions": table_captions,
        "tables": tables[:40],
    }


def main() -> int:
    ref = ep1581_reference_docx()
    if not ref.exists():
        print(f"Reference not found: {ref}")
        return 1
    payload = parse_reference_docx(ref)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(yaml.dump(payload, sort_keys=False, allow_unicode=True, width=120), encoding="utf-8")
    print(f"Wrote {OUT}")
    print(f"  headings={payload['stats']['headings']} tables={payload['stats']['tables']} figures={payload['stats']['figure_captions']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
