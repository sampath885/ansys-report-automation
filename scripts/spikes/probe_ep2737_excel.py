"""Probe EP2737 Excel layout for Phase 4 adapter."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

from ansys_report.ep2737_paths import ep2737_case_root

ROOT = ep2737_case_root()
WB = ROOT / "OLF Mechanical Calculation of BOM IDs EP2737_1.xlsx"


def dump_sheet(ws, max_row: int = 80) -> None:
    for i, row in enumerate(ws.iter_rows(max_row=max_row, values_only=True), start=1):
        if any(c is not None for c in row):
            safe = tuple("" if c is None else str(c).encode("ascii", "replace").decode() for c in row[:12])
            print(f"{i:3} {safe}")


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    wb = load_workbook(WB, data_only=True)
    for sn in wb.sheetnames:
        print(f"\n===== {sn!r} =====")
        dump_sheet(wb[sn], max_row=100 if "Flange" in sn else 60)
    wb.close()


if __name__ == "__main__":
    main()
