"""Phase 0 inventory: parse EP2737 reference report + data sources → YAML spec."""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import yaml
from docx import Document

ROOT = Path(__file__).resolve().parents[1] / "EP 2737"
WB = ROOT / "Structural Analysis_EP2737"
DP0 = WB / "EP2737_files" / "dp0"
OUT = Path(__file__).resolve().parents[1] / "docs" / "ep2737_data_spec.yaml"


def parse_reference_docx() -> dict:
    doc = Document(ROOT / "Design Report_EP2737_UPDATED.docx")
    sections = [
        {"level": p.style.name, "title": p.text.strip()}
        for p in doc.paragraphs
        if p.style.name.startswith("Heading") and p.text.strip()
    ]
    captions = [
        p.text.strip()
        for p in doc.paragraphs
        if p.style.name == "Caption" and p.text.strip()
    ]
    figures = [c for c in captions if c.lower().startswith("figure")]
    table_captions = [c for c in captions if c.lower().startswith("table")]

    tables = []
    for ti, table in enumerate(doc.tables):
        first_row = [c.text.strip()[:80] for c in table.rows[0].cells] if table.rows else []
        tables.append(
            {
                "index": ti,
                "rows": len(table.rows),
                "cols": len(table.columns) if table.rows else 0,
                "first_row": first_row,
            }
        )

    return {
        "path": "Design Report_EP2737_UPDATED.docx",
        "stats": {
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "figure_captions": len(figures),
            "table_captions": len(table_captions),
            "headings": len(sections),
        },
        "cover": {
            "title": "WELDED PLANE PIPE FLANGE",
            "bom_id": "EP 2737 – S5",
            "customer": "M/s ULTRA DIMENSIONS PVT. LTD., VISAKHAPATNAM",
        },
        "sections": sections,
        "figure_captions": figures,
        "table_captions": table_captions,
        "tables": tables,
    }


def suggest_export_path(caption: str) -> str | None:
    c = caption.lower()
    if "3d model" in c or "sectional view" in c:
        return "geometry/cad_iso.png"
    if "2d drawing" in c:
        return "geometry/cad_section.png"
    if "welded plane pipe flange" in c and "figure 4" in c:
        return "geometry/flange_drawing.png"
    if "model orientation" in c:
        return "geometry/model_orientation.png"
    if c.startswith("figure 6:") or (c.startswith("figure") and ": mesh" in c):
        return "mesh/mesh_global.png"
    if "aspect ratio" in c:
        return "mesh/aspect_ratio.png"
    if "skewness" in c:
        return "mesh/skewness.png"
    if "jacobian" in c:
        return "mesh/jacobian.png"
    if "mesh quality" in c:
        return "mesh/mesh_quality.png"
    if "corner angle" in c:
        return "mesh/max_corner_angle.png"
    if "beam" in c:
        return "modelling/bolt_beam.png"
    if "contact" in c:
        return "modelling/contacts.png"
    if "earth gravity" in c:
        return "static/earth_gravity.png"
    if "pressure" in c:
        return "static/pressure.png"
    if "excitation" in c and "x" in c:
        return "harmonic_x/excitation.png"
    if "excitation" in c and "y" in c:
        return "harmonic_y/excitation.png"
    if "excitation" in c and "z" in c:
        return "harmonic_z/excitation.png"
    if "mode shape" in c:
        m = re.search(r"mode\s*(\d+)", c)
        n = m.group(1) if m else "N"
        return f"modal/mode{n}.png"
    if "total deformation" in c or "total deform" in c:
        return "static/total_deformation.png"
    if "equivalent" in c and "stress" in c:
        return "static/vonmises.png"
    if "fixed support" in c:
        return "static/fixed_support.png"
    if "displacement" in c and "harmonic" in c:
        return "harmonic_x/amplitude.png"
    if "von" in c or "stress" in c:
        return "static/vonmises.png"
    return None


def parse_wbpj_systems() -> list[dict]:
    text = (WB / "EP2737.wbpj").read_text(encoding="utf-8", errors="replace")
    seen: set[str] = set()
    systems: list[dict] = []
    for m in re.finditer(r'"UniqueSystemDirectoryName": "(SYS(?:-\d+)?)"', text):
        folder = m.group(1)
        if folder in seen:
            continue
        seen.add(folder)
        chunk = text[max(0, m.start() - 2000) : m.end() + 100]
        dn = re.findall(r'"DisplayText": "([^"]+)"', chunk)
        display_name = dn[-1] if dn else folder
        at = re.search(r'"AnalysisType": "(\w+)"', chunk)
        systems.append(
            {
                "folder": folder,
                "display_name": display_name,
                "workbench_analysis_type": at.group(1) if at else "unknown",
            }
        )
    return sorted(
        systems,
        key=lambda s: (0 if s["folder"] == "SYS" else int(s["folder"].split("-")[1])),
    )


REPORT_SYSTEM_KEY = {
    "SYS": "static_structural",
    "SYS-1": "modal",
    "SYS-2": "vibration_x",
    "SYS-3": "vibration_y",
    "SYS-4": "vibration_z",
    "SYS-5": "shock_plus_x",
    "SYS-6": "shock_plus_y",
    "SYS-7": "shock_plus_z",
    "SYS-8": "shock_minus_x",
    "SYS-9": "shock_minus_y",
    "SYS-10": "shock_minus_z",
}


def mech_inventory(folder: str) -> dict:
    mech = DP0 / folder / "MECH"
    info: dict = {"mech_relative": str(mech.relative_to(ROOT)).replace("\\", "/")}
    if not mech.exists():
        info["status"] = "missing"
        return info

    rst_files = sorted(p.name for p in mech.glob("file*.rst"))
    info["rst_files"] = rst_files
    info["primary_rst"] = "file.rst" if (mech / "file.rst").exists() else None
    info["has_mcf"] = (mech / "file.mcf").exists()
    info["has_mode_files"] = bool(list(mech.glob("*.mode")))

    caerep_path = mech / "CAERep.xml"
    if caerep_path.exists():
        text = caerep_path.read_text(encoding="utf-8", errors="replace")
        pm = re.search(r'<ProjectName PropType="string">([^<]+)</ProjectName>', text)
        at = re.search(r'<AnalysisType PropType="string">(\w+)</AnalysisType>', text)
        info["project_name"] = pm.group(1) if pm else None
        info["solver_analysis_type"] = at.group(1) if at else None

    ds_path = mech / "ds.dat"
    if ds_path.exists():
        dst = ds_path.read_text(encoding="utf-8", errors="replace")[:100_000]
        if "antype,harm" in dst:
            info["antype"] = "harmonic"
        elif "antype,modal" in dst or (mech / "file.db").exists():
            info["antype"] = "modal"
        else:
            info["antype"] = "static"

    solve_path = mech / "solve.out"
    if solve_path.exists():
        sout = solve_path.read_text(encoding="utf-8", errors="replace")
        if re.search(r"NUMBER OF ERROR MESSAGES\s+=\s+0", sout):
            info["mapdl_errors"] = 0
        elif "RUN COMPLETED" in sout:
            info["mapdl_errors"] = 0
        else:
            info["mapdl_errors"] = "unknown"
    return info


def parse_static_metadata() -> dict:
    caerep = (DP0 / "SYS" / "MECH" / "CAERep.xml").read_text(encoding="utf-8", errors="replace")
    steps = [float(t) for t in re.findall(r"<EndTime>([\d.]+)</EndTime>", caerep)]
    pretensions = [
        float(v) for v in re.findall(r'<Preload PropType="double"[^>]*>([\d.]+)</Preload>', caerep)
    ]

    bodies = []
    for m in re.finditer(
        r'<BodyAttributes[^>]*>.*?<Caption PropType="string">([^<]+)</Caption>'
        r'.*?<Mass PropType="double"[^>]*>([\d.E+-]+)</Mass>'
        r'.*?<MaterialRep[^>]*>.*?<MaterialName PropType="string">([^<]+)</MaterialName>',
        caerep,
        re.S,
    ):
        bodies.append(
            {
                "name": m.group(1),
                "mass_tonne": float(m.group(2)),
                "mass_kg": round(float(m.group(2)) * 1000, 4),
                "material": m.group(3),
            }
        )

    cog_m = re.search(
        r"<Centroid ObjId=\d+ Type=\"DSGePoint3d\".*?"
        r'<Coordinates PropType="vector&lt;double>"[^>]*>([^<]+)</Coordinates>',
        caerep,
        re.S,
    )
    cog = None
    if cog_m:
        parts = [float(x.strip()) for x in cog_m.group(1).split(",")]
        cog = {"x_mm": parts[0], "y_mm": parts[1], "z_mm": parts[2]}

    contacts = re.findall(r'<Caption PropType="string">([^<]*Contact[^<]*)</Caption>', caerep)

    return {
        "load_steps": steps,
        "report_result_step": 3,
        "report_result_step_note": "Step 3 applies full pressure after pretension (steps 1-2); confirm against reference report.",
        "bolt_preload_n": pretensions[0] if pretensions else None,
        "bolt_count": len(pretensions),
        "bodies": bodies,
        "assembly_cog_mm": cog,
        "contacts": contacts[:5],
    }


def parse_modal_frequencies() -> list[dict]:
    solve = (DP0 / "SYS-1" / "MECH" / "solve.out").read_text(encoding="utf-8", errors="replace")
    freqs: list[dict] = []
    in_block = False
    for line in solve.splitlines():
        if "FREQUENCIES FROM BLOCK LANCZOS" in line:
            in_block = True
            continue
        if in_block:
            m = re.match(r"\s*(\d+)\s+([\d.E+-]+)", line)
            if m:
                freqs.append({"mode": int(m.group(1)), "freq_hz": float(m.group(2))})
            elif freqs and "ELEMENT RESULT" in line:
                break
    return freqs


def parse_mesh_stats() -> dict:
    ds_path = DP0 / "SYS" / "MECH" / "ds.dat"
    ds = ds_path.read_text(encoding="utf-8", errors="replace")[:500_000]
    node_count = None
    elem_count = None
    for pat in [
        r"Number of nodes\s*=\s*(\d+)",
        r"/COM,\s*(\d+)\s+nodes",
    ]:
        m = re.search(pat, ds, re.I)
        if m:
            node_count = int(m.group(1))
            break
    m2 = re.search(r"Number of elements\s*=\s*(\d+)", ds, re.I)
    if m2:
        elem_count = int(m2.group(1))
    return {
        "node_count": node_count or 323537,
        "element_count": elem_count,
        "note": "node_count defaults to 323537 from CAERep scan if not found in ds.dat header",
    }


def parse_excel() -> dict:
    result = {}
    for fname in [
        "OLF Mechanical Calculation of BOM IDs EP2737_1.xlsx",
        "Bolt pre load.xlsx",
    ]:
        path = ROOT / fname
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = {}
        for sn in wb.sheetnames:
            ws = wb[sn]
            sample = []
            for row in ws.iter_rows(max_row=15, values_only=True):
                if any(c is not None for c in row):
                    sample.append([str(c)[:100] if c is not None else None for c in row[:8]])
            sheets[sn] = {"sample_rows": sample[:12]}
        result[fname] = {
            "relative_path": fname,
            "sheets": sheets,
        }
    return result


def build_spec() -> dict:
    ref = parse_reference_docx()
    figure_slots = [
        {
            "caption": cap,
            "suggested_export_path": suggest_export_path(cap),
        }
        for cap in ref["figure_captions"]
    ]

    systems = {}
    for s in parse_wbpj_systems():
        key = REPORT_SYSTEM_KEY.get(s["folder"], s["folder"])
        systems[key] = {
            **s,
            "report_section_key": key,
            **mech_inventory(s["folder"]),
        }

    png_count = len(list(ROOT.rglob("*.png"))) + len(list(ROOT.rglob("*.jpg")))

    return {
        "meta": {
            "phase": 0,
            "bom_id": "EP 2737",
            "purpose": "Ground-truth data contract for report automation",
            "regenerate": "python scripts/inventory_ep2737.py",
        },
        "reference_template": {
            k: v for k, v in ref.items() if k not in ("figure_captions",)
        },
        "figure_export_slots": figure_slots,
        "report_to_ansys_mapping": [
            {
                "report_section": "Static Structural Analysis",
                "system_key": "static_structural",
                "workbench_folder": "SYS",
            },
            {
                "report_section": "Modal Analysis",
                "system_key": "modal",
                "workbench_folder": "SYS-1",
            },
            {
                "report_section": "Vibration Resistance Analysis (X/Y/Z)",
                "system_key": "vibration_x | vibration_y | vibration_z",
                "workbench_folder": "SYS-2 | SYS-3 | SYS-4",
            },
            {
                "report_section": "Shock Analysis - Equivalent Static Approach",
                "system_key": "shock_plus_x … shock_minus_z",
                "workbench_folder": "SYS-5 … SYS-10",
            },
            {
                "report_section": "MECHANICAL DESIGN CALCULATION",
                "system_key": "excel",
                "workbench_folder": "N/A",
            },
        ],
        "paths": {
            "case_root": "EP 2737",
            "workbench_project": "EP 2737/Structural Analysis_EP2737",
            "wbpj": "EP 2737/Structural Analysis_EP2737/EP2737.wbpj",
            "cad_step": "EP 2737/11-EP2737.STEP",
            "geometry_scdocx": "EP 2737/Structural Analysis_EP2737/EP2737_files/dp0/SYS/DM/SYS.scdocx",
            "exports": "EP 2737/exports",
            "excel_design_calcs": "EP 2737/OLF Mechanical Calculation of BOM IDs EP2737_1.xlsx",
            "excel_bolt_preload": "EP 2737/Bolt pre load.xlsx",
        },
        "ansys_environment": {
            "version": "2025 R2",
            "version_build": "25.2.1.0",
            "solver_units": "N, mm, MPa (solver unit code 6)",
            "pyansys_note": "Pin ansys-dpf-core to build matching ANSYS 2025 R2 on target laptop",
        },
        "ansys_systems": systems,
        "static_analysis": parse_static_metadata(),
        "modal_analysis": {
            "modes_requested": 20,
            "modes_for_report": 6,
            "frequencies_hz": parse_modal_frequencies(),
            "golden_mode1_hz": 421.1241433955,
        },
        "mesh": parse_mesh_stats(),
        "excel": parse_excel(),
        "excel_report_tables": [
            {
                "report_table": "Table 25: Flange Thickness calculation",
                "workbook": "OLF Mechanical Calculation of BOM IDs EP2737_1.xlsx",
                "sheet": "Welded Flange Thickness Calcula",
            },
            {
                "report_table": "Table 26: Bolt pretension calculations",
                "workbook": "OLF Mechanical Calculation of BOM IDs EP2737_1.xlsx",
                "sheet": "Bolt pretension load with ",
            },
            {
                "report_table": "Table 27: Fatigue Strength Calculations",
                "workbook": "OLF Mechanical Calculation of BOM IDs EP2737_1.xlsx",
                "sheet": "bolt pre load",
            },
        ],
        "images": {
            "exports_root": "EP 2737/exports",
            "status": "present" if png_count else "missing",
            "png_jpg_count": png_count,
        },
        "known_issues": [
            "Harmonic SYS-2/3/4: Workbench Solution cell error; MAPDL solve.out reports 0 errors",
            "Shock SYS-5…10: Workbench Setup refresh error; ds.dat and file.rst present",
            "Project originally on \\\\sambaserver\\FE_ANALYSIS\\...; use local paths",
            "IMPLEMENTATION.md references EP1763 Excel sheet names; EP2737 uses different sheet names",
        ],
        "phase_0_exit_checklist": [
            "Reference DOCX section map matches ansys_systems keys",
            "Modal mode-1 frequency 421.124 Hz documented as DPF spike golden value",
            "Static load step 3 identified for final stress/deformation (engineer confirm)",
            "Figure export slots defined for all 105 figure captions",
            "Excel sheet mapping documented for design calc tables",
        ],
    }


def main() -> None:
    spec = build_spec()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8") as f:
        yaml.dump(spec, f, default_flow_style=False, allow_unicode=True, sort_keys=False, width=100)
    print(f"Wrote {OUT}")
    print(f"  Systems: {len(spec['ansys_systems'])}")
    print(f"  Figure slots: {len(spec['figure_export_slots'])}")
    print(f"  Modal modes: {len(spec['modal_analysis']['frequencies_hz'])}")


if __name__ == "__main__":
    main()
