"""Phase 7 tests — harmonic Y/Z, shock, narratives (no images)."""

import json
import os
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
GOLDEN = Path(__file__).resolve().parent / "fixtures" / "ep2737_golden"
CONFIG = REPO / "config" / "project.ep2737.yaml"
WB = REPO / "EP 2737" / "Structural Analysis_EP2737" / "EP2737_files" / "dp0"

pytestmark_harmonic = pytest.mark.skipif(
    os.getenv("ANSYS_AVAILABLE") != "1",
    reason="Set ANSYS_AVAILABLE=1 for harmonic Y/Z DPF tests",
)


@pytest.fixture
def ep2737_cfg():
    if not (REPO / "EP 2737").exists():
        pytest.skip("EP2737 not in workspace")
    from ansys_report.config import load_project_config

    return load_project_config(CONFIG)


@pytestmark_harmonic
def test_harmonic_y_golden():
    from ansys_report.extract.dpf_harmonic import extract_harmonic_peak

    golden = json.loads((GOLDEN / "2e_harmonic_y.json").read_text(encoding="utf-8"))
    result = extract_harmonic_peak(WB / "SYS-3" / "MECH" / "file.rst")
    assert result.peak_displacement_mm == pytest.approx(golden["peak_displacement_mm"], abs=0.05)


@pytestmark_harmonic
def test_harmonic_z_golden():
    from ansys_report.extract.dpf_harmonic import extract_harmonic_peak

    golden = json.loads((GOLDEN / "2f_harmonic_z.json").read_text(encoding="utf-8"))
    result = extract_harmonic_peak(WB / "SYS-4" / "MECH" / "file.rst")
    assert result.peak_displacement_mm == pytest.approx(golden["peak_displacement_mm"], abs=0.05)


@pytest.mark.skipif(
    os.getenv("ANSYS_AVAILABLE") != "1",
    reason="Set ANSYS_AVAILABLE=1 for live shock DPF golden test",
)
def test_shock_all_golden(ep2737_cfg):
    from ansys_report.config import load_project_config
    from ansys_report.sections.shock import ShockSection
    from ansys_report.scanner import scan_project

    golden = json.loads((GOLDEN / "7_shock_all.json").read_text(encoding="utf-8"))
    cfg = load_project_config(CONFIG)
    inv = scan_project(
        cfg.project_dir,
        cfg.image_folder,
        cfg.excel_calcs,
        case_root=cfg.case_root,
        excel_bolt_preload=cfg.excel_bolt_preload,
    )
    data = ShockSection().extract(inv, cfg)
    assert len(data["directions"]) == 6
    for got, exp in zip(data["directions"], golden["directions"], strict=True):
        if got.get("max_stress_mpa") is None:
            pytest.skip("DPF unavailable for shock extraction")
        assert got["max_stress_mpa"] == pytest.approx(exp["max_stress_mpa"], abs=0.5)


def test_harmonic_narrative_caution():
    from ansys_report.config import ThresholdsConfig
    from ansys_report.models import HarmonicPeakResult
    from ansys_report.narrative import rules

    cfg = type("C", (), {"operating_freq_hz": [10.0, 200.0], "primary_yield_mpa": 170.0})()
    peak = HarmonicPeakResult(peak_displacement_mm=11.38, peak_frequency_hz=3.91, num_frequency_sets=15)
    narr = rules.narrate_harmonic(peak, "X", cfg, ThresholdsConfig())
    assert narr.verdict == "CAUTION"


@pytest.mark.skipif(
    os.getenv("ANSYS_AVAILABLE") != "1",
    reason="Set ANSYS_AVAILABLE=1 for Phase 7 DPF spike runner (slow)",
)
def test_phase7_spike_runner():
    import importlib.util

    script = REPO / "scripts" / "spikes" / "run_ep2737_phase7.py"
    spec = importlib.util.spec_from_file_location("run_ep2737_phase7", script)
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(mod)
    assert mod.run_phase7(compare=True) == 0


def test_report_includes_vibration_and_shock(ep2737_cfg, tmp_path):
    from ansys_report.report.ep2737_overlay import apply_golden_dpf_overlay
    from ansys_report.report.render import render_report
    from docx import Document

    ctx = {
        "bom_id": ep2737_cfg.bom_id,
        "title": ep2737_cfg.title,
        "customer": ep2737_cfg.customer,
        "prepared_by": ep2737_cfg.prepared_by.model_dump(),
        "checked_by": ep2737_cfg.checked_by.model_dump(),
        "approved_by": ep2737_cfg.approved_by.model_dump(),
        "equipment": {"bom_id": ep2737_cfg.bom_id, "title": ep2737_cfg.title, "bodies": [], "assembly": None},
        "modelling": {"node_count": 294165, "element_count": 202546, "load_steps": [1, 2, 3], "contacts": []},
        "design_calcs": {"end_flange": [], "effort": []},
        "narrative": {"executive_summary": "Phase 7 test"},
    }
    apply_golden_dpf_overlay(ctx, ep2737_cfg, force=True)

    tpl = ep2737_cfg.template_path or REPO / "templates" / "EP2737_report_template.docx"
    out = tmp_path / "phase7.docx"
    render_report(tpl, ctx, out)

    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            text += "\n" + " ".join(c.text for c in row.cells)

    assert "Vibration Resistance" in text
    assert "Shock Analysis" in text
    assert "+X" in text
    assert "255" in text
