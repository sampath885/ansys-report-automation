"""Phase 6 render tests — first EP2737 DOCX without images."""

import os
import zipfile
from pathlib import Path

import pytest
from docx import Document

REPO = Path(__file__).resolve().parents[1]
CONFIG = REPO / "config" / "project.ep2737.yaml"
GOLDEN = Path(__file__).resolve().parent / "fixtures" / "ep2737_golden"
TEMPLATE = REPO / "templates" / "EP2737_report_template.docx"


@pytest.fixture(scope="session")
def ep2737_template(repo_root) -> Path:
    tpl = repo_root / "templates" / "EP2737_report_template.docx"
    if not tpl.exists():
        from ansys_report.report.template_builder import create_ep2737_template

        create_ep2737_template(tpl)
    return tpl


@pytest.fixture
def ep2737_case():
    if not (REPO / "EP 2737" / "Structural Analysis_EP2737").exists():
        pytest.skip("EP2737 project not in workspace")


@pytest.fixture
def ep2737_cfg(ep2737_case):
    from ansys_report.config import load_project_config

    return load_project_config(CONFIG)


def _doc_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_render_ep2737_from_inventory(ep2737_cfg, ep2737_template, tmp_path):
    from ansys_report.report.context_builder import build_context
    from ansys_report.report.ep2737_overlay import apply_golden_dpf_overlay
    from ansys_report.report.render import render_report
    from ansys_report.scanner import scan_project

    inventory = scan_project(
        ep2737_cfg.project_dir,
        ep2737_cfg.image_folder,
        ep2737_cfg.excel_calcs,
        case_root=ep2737_cfg.case_root,
        excel_bolt_preload=ep2737_cfg.excel_bolt_preload,
    )
    dpf_keys = {"modal", "static", "harmonic_x", "harmonic_y", "harmonic_z", "shock"}
    saved = list(ep2737_cfg.sections_enabled)
    ep2737_cfg.sections_enabled = [s for s in saved if s not in dpf_keys]
    ctx, _ = build_context(ep2737_cfg, inventory=inventory, use_ai=False)
    ep2737_cfg.sections_enabled = saved
    apply_golden_dpf_overlay(ctx, ep2737_cfg, force=True)

    out = tmp_path / "EP_2737_report.docx"
    render_report(ep2737_template, ctx, out)

    text = _doc_text(out)
    assert "EP 2737" in text
    assert "WELDED PLANE PIPE FLANGE" in text
    assert "421.12" in text
    assert "239.6" in text or "239.63" in text
    assert "11832" in text
    assert "294165" in text
    assert "Vibration Resistance" in text
    assert "Shock Analysis" in text
    assert "FACTOR OF SAFETY" in text.upper() or "1.8" in text
    # Must open in Word — no legacy stylesWithEffects part
    with zipfile.ZipFile(out) as z:
        assert "word/stylesWithEffects.xml" not in z.namelist()


def test_golden_overlay_only_when_needed(ep2737_cfg):
    from ansys_report.report.ep2737_overlay import apply_golden_dpf_overlay, needs_golden_overlay

    ctx = {"modal": {"modes": []}, "static": {"max_stress_mpa": None}}
    assert needs_golden_overlay(ctx)
    assert apply_golden_dpf_overlay(ctx, ep2737_cfg, golden_dir=GOLDEN)
    assert ctx["modal"]["modes"][0]["freq_hz"] == pytest.approx(421.124, abs=0.01)
    assert ctx["static"]["max_stress_mpa"] == pytest.approx(239.63, abs=0.1)


def test_build_spike():
    import importlib.util

    script = REPO / "scripts" / "spikes" / "run_ep2737_build.py"
    spec = importlib.util.spec_from_file_location("run_ep2737_build", script)
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(mod)
    assert mod.run_build(golden_dpf=True, out=REPO / "output" / "test_ep2737_build.docx") == 0

@pytest.mark.skipif(
    os.getenv("ANSYS_AVAILABLE") != "1",
    reason="Set ANSYS_AVAILABLE=1 for live DPF build test",
)
def test_render_ep2737_live_dpf(ep2737_cfg, ep2737_template, tmp_path):
    from ansys_report.report.context_builder import build_context
    from ansys_report.report.render import render_report
    from ansys_report.scanner import scan_project

    inventory = scan_project(
        ep2737_cfg.project_dir,
        ep2737_cfg.image_folder,
        ep2737_cfg.excel_calcs,
        case_root=ep2737_cfg.case_root,
        excel_bolt_preload=ep2737_cfg.excel_bolt_preload,
    )
    ctx, _ = build_context(ep2737_cfg, inventory=inventory, use_ai=False)
    assert ctx["modal"]["modes"], "Expected live modal extraction"
    assert ctx["static"]["max_stress_mpa"] is not None

    out = tmp_path / "EP_2737_live.docx"
    render_report(ep2737_template, ctx, out)
    assert out.exists()


@pytest.mark.skipif(
    os.getenv("RUN_WORD_COM") != "1",
    reason="Set RUN_WORD_COM=1 to run Word desktop open test",
)
def test_word_opens_rendered_report(ep2737_cfg, ep2737_template, tmp_path):
    """Microsoft Word must open the generated DOCX on Windows."""
    pytest.importorskip("win32com")
    from ansys_report.report.context_builder import build_context
    from ansys_report.report.ep2737_overlay import apply_golden_dpf_overlay
    from ansys_report.report.render import render_report
    from ansys_report.scanner import scan_project

    inventory = scan_project(
        ep2737_cfg.project_dir,
        ep2737_cfg.image_folder,
        ep2737_cfg.excel_calcs,
        case_root=ep2737_cfg.case_root,
        excel_bolt_preload=ep2737_cfg.excel_bolt_preload,
    )
    ctx, _ = build_context(ep2737_cfg, inventory=inventory, use_ai=False)
    apply_golden_dpf_overlay(ctx, ep2737_cfg)
    out = tmp_path / "word_compat.docx"
    render_report(ep2737_template, ctx, out)

    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(str(out.resolve()))
        doc.Close(False)
    finally:
        word.Quit()
