"""Tests for EP1581 section content matrix and block-based renderer."""

from __future__ import annotations

from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
SPEC = REPO / "config" / "ep1581_section_content.yaml"
CONFIG = REPO / "config" / "project.ep2741.yaml"


def test_load_ep1581_section_content_spec():
    from ansys_report.report.section_spec import load_section_content_spec

    spec = load_section_content_spec(SPEC)
    assert spec.layout == "ep1581"
    assert spec.content_start_marker == "Revision Log"
    keys = [s.key for s in spec.sections]
    assert "vibration_intro" in keys
    assert "modelling" in keys
    assert "loads" not in keys


def test_ep1581_modelling_only_mesh_figure():
    from ansys_report.report.section_spec import load_section_content_spec

    spec = load_section_content_spec(SPEC)
    modelling = next(s for s in spec.sections if s.key == "modelling")
    figure_slots = [b.slot for b in modelling.blocks if b.type == "figure"]
    assert figure_slots == ["mesh_global"]


def test_ep1581_design_calcs_use_figure_placeholder():
    from ansys_report.report.section_spec import load_section_content_spec

    spec = load_section_content_spec(SPEC)
    design = next(s for s in spec.sections if s.key == "design_calcs")
    repeat = next(b for b in design.blocks if b.type == "repeat")
    fig_blocks = [b for b in repeat.blocks if b.type in ("figure", "figure_placeholder")]
    assert len(fig_blocks) == 1
    assert fig_blocks[0].type == "figure_placeholder"
    assert fig_blocks[0].when_field == "raw_rows"


def test_ep1581_modelling_has_no_joints_section():
    from ansys_report.report.section_spec import load_section_content_spec

    spec = load_section_content_spec(SPEC)
    modelling = next(s for s in spec.sections if s.key == "modelling")
    headings = [b.text for b in modelling.blocks if b.type == "heading"]
    assert "Modelling approach for joints" not in headings


def test_ep1581_modal_under_loads_parent():
    from ansys_report.report.section_spec import load_section_content_spec

    spec = load_section_content_spec(SPEC)
    modal = next(s for s in spec.sections if s.key == "modal")
    assert modal.parent_section == "Details of the Loading and BC"
    assert modal.parent_heading_level == 1


def test_assemble_modelling_includes_mesh_figure(tmp_path, tiny_png):
    import shutil

    from ansys_report.config import load_project_config
    from ansys_report.report.block_assembler import assemble_report_document

    exports = tmp_path / "exports"
    mesh_dir = exports / "mesh"
    mesh_dir.mkdir(parents=True)
    shutil.copy(tiny_png, mesh_dir / "mesh.png")

    cfg = load_project_config(CONFIG)
    cfg.sections_enabled = ["modelling"]
    ctx = {
        "modelling": {"node_count": 100, "element_count": 50},
        "images": {"mesh_global": mesh_dir / "mesh.png"},
        "image_root": str(exports),
        "methodology": {},
        "reference_tables": {
            "mesh_control_reference": {"headers": ["A"], "rows": [["1"]]},
            "mesh_quality": {"headers": ["A"], "rows": [["1"]]},
        },
    }
    doc = assemble_report_document(ctx, cfg)
    modelling = next(s for s in doc.sections if s.key == "modelling")
    mesh_figs = [b for b in modelling.blocks if b.kind == "figure" and b.image_path]
    assert len(mesh_figs) == 1
    assert "Figure 8" in (mesh_figs[0].caption or "")


def test_assemble_ep1581_document_minimal():
    from ansys_report.config import load_project_config
    from ansys_report.report.block_assembler import assemble_report_document
    from ansys_report.report.table_builders import enrich_context_tables

    cfg = load_project_config(CONFIG)
    cfg.sections_enabled = ["modal", "static"]
    ctx = {
        "modal": {
            "modes": [{"index": 1, "freq_hz": 100.0}],
            "intro_text": "Modal intro.",
            "summary_table": [{"sr_no": 1, "freq_hz": 100.0, "operating_frequency": "10-200 Hz", "dominant_direction": "—", "remark": "OK"}],
            "narrative": {"conclusions": ["Modal OK."]},
        },
        "static": {
            "max_stress_mpa": 50.0,
            "conclusion_table": [{"sr_no": 1, "material": "Steel", "location": "Body", "stress_mpa": 50.0, "allowable_mpa": 100.0, "remarks": "OK"}],
            "narrative": {"conclusions": ["Static OK."]},
        },
        "methodology": {
            "modal_methodology": "Modal method.",
            "working_principle": "Valve works.",
        },
    }
    enrich_context_tables(ctx, cfg)
    doc = assemble_report_document(ctx, cfg)
    modal = next(s for s in doc.sections if s.key == "modal")
    assert any(b.kind == "heading" and b.text == "Details of the Loading and BC" for b in modal.blocks)
    assert any(b.kind == "table" for b in modal.blocks)


def test_ep1581_block_render_caption_after_image(tmp_path):
    from ansys_report.config import load_project_config
    from ansys_report.report.block_assembler import assemble_report_document
    from ansys_report.report.block_render import render_blocks_document
    from ansys_report.report.blocks import RenderBlock, RenderDocument, RenderSection
    from docx import Document

    cfg = load_project_config(CONFIG)
    shell = REPO / "templates" / "EP1581_style_shell.docx"
    if not shell.exists():
        pytest.skip("EP1581 style shell not generated")

    png = tmp_path / "fig.png"
    try:
        from PIL import Image

        Image.new("RGB", (100, 100), color="red").save(png)
    except ImportError:
        pytest.skip("Pillow not available")

    doc_model = RenderDocument(
        sections=[
            RenderSection(
                key="test",
                blocks=[
                    RenderBlock(kind="figure", caption="Figure 1 - Test", image_path=str(png), slot="test"),
                ],
            )
        ]
    )
    out = tmp_path / "out.docx"
    render_blocks_document(doc_model, {}, out, style_shell_path=shell, layout="ep1581")
    saved = Document(str(out))
    assert out.is_file()
    caption_paras = [p for p in saved.paragraphs if p.text.strip() == "Figure 1 - Test"]
    assert caption_paras, "Expected caption paragraph"
    assert caption_paras[0].style.name == "Caption"


def test_ep1581_optional_missing_figure_omitted():
    from ansys_report.config import load_project_config
    from ansys_report.report.block_assembler import assemble_report_document

    cfg = load_project_config(CONFIG)
    cfg.sections_enabled = ["equipment"]
    ctx = {
        "equipment": {},
        "methodology": {"working_principle": "Valve works."},
        "reference_tables": {},
    }
    doc = assemble_report_document(ctx, cfg)
    equipment = next(s for s in doc.sections if s.key == "equipment")
    assert not any(b.kind == "pending" for b in equipment.blocks)
    assert not any(b.slot == "cad_isometric" for b in equipment.blocks)


def test_project_ep2741_uses_ep1581_layout():
    from ansys_report.config import load_project_config

    cfg = load_project_config(CONFIG)
    assert cfg.layout == "ep1581"
    assert cfg.strict_mode is True
    assert "loads" not in cfg.sections_enabled
    assert cfg.section_content_path.name == "ep1581_section_content.yaml"
