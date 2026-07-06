"""Tests for EP1581 DOCX formatting helpers."""

from __future__ import annotations

from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]


def test_ep1581_table_full_width_bold_headers(tmp_path):
    from docx import Document
    from ansys_report.report.ep1581_docx_format import add_table

    doc = Document()
    add_table(doc, ["Col A", "Col B"], [["1", "2"]])
    out = tmp_path / "t.docx"
    doc.save(str(out))

    import zipfile
    from xml.etree import ElementTree as ET

    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(out) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    tbl = root.find(f".//{W}tbl")
    assert tbl is not None
    tbl_w = tbl.find(f"{W}tblPr/{W}tblW")
    assert tbl_w is not None
    assert tbl_w.get(f"{W}w") == "5000"
    assert tbl_w.get(f"{W}type") == "pct"
    bold = tbl.findall(f".//{W}tr[1]//{W}b")
    assert bold, "Expected bold header cells"


def test_normalize_docx_portrait_a4_on_landscape_template(tmp_path):
    from docx import Document

    from ansys_report.report.ep1581_docx_format import add_section_break, normalize_docx_portrait_a4

    doc = Document()
    add_section_break(doc, "landscape")
    raw = tmp_path / "landscape.docx"
    doc.save(str(raw))
    normalize_docx_portrait_a4(raw)

    import zipfile
    from xml.etree import ElementTree as ET

    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(raw) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    sizes = [(pg.get(f"{W}w"), pg.get(f"{W}h"), pg.get(f"{W}orient")) for pg in root.iter(f"{W}pgSz")]
    assert sizes
    assert all(w == "11906" and h == "16838" and orient is None for w, h, orient in sizes)


def test_regenerate_ep1581_templates_from_reference():
    ref = Path.home() / "Downloads" / "EP1581_29APR_R0 (1).docx"
    if not ref.exists():
        pytest.skip("EP1581 reference DOCX not in Downloads")
    from scripts.export_ep1581_templates import export_layout_template, export_style_shell

    layout = export_layout_template(reference_docx=ref)
    shell = export_style_shell(reference_docx=ref)
    assert layout.exists()
    assert shell.exists()
