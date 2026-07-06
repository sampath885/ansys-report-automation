"""Tests for EP1581 figure centering and descriptions in block render."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from ansys_report.report.block_render import _render_block
from ansys_report.report.blocks import RenderBlock


def test_ep1581_figure_centered_with_description(tmp_path, tiny_png):
    doc = Document()
    block = RenderBlock(
        kind="figure",
        image_path=str(tiny_png),
        caption="Figure 1 - Test",
        description="The computed maximum stress of 10.0 MPa is within limits.",
    )
    _render_block(doc, block, {}, ep1581=True)

    assert len(doc.paragraphs) >= 2
    pic_para = doc.paragraphs[0]
    assert pic_para.alignment == WD_PARAGRAPH_ALIGNMENT.CENTER
    assert len(doc.inline_shapes) == 1
    desc_found = any("10.0 MPa" in p.text for p in doc.paragraphs)
    assert desc_found
