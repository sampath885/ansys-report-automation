"""Tests for EP2737 section content matrix and block-based renderer."""

from __future__ import annotations

from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
SPEC = REPO / "config" / "ep2737_section_content.yaml"
CONFIG = REPO / "config" / "project.ep2737.yaml"


def test_load_section_content_spec():
    from ansys_report.report.section_spec import load_section_content_spec

    spec = load_section_content_spec(SPEC)
    assert spec.version == 1
    keys = [s.key for s in spec.sections]
    assert "harmonic_x" in keys
    assert "shock" in keys
    assert "design_calcs" in keys


def test_harmonic_spec_has_figures_not_tables():
    from ansys_report.report.section_spec import load_section_content_spec

    spec = load_section_content_spec(SPEC)
    hx = next(s for s in spec.sections if s.key == "harmonic_x")
    block_types = [b.type for b in hx.blocks]
    assert "figure" in block_types
    assert "table" not in block_types
    assert any(b.type == "scalar" for b in hx.blocks)


def test_assemble_document_from_minimal_context():
    from ansys_report.config import load_project_config
    from ansys_report.report.block_assembler import assemble_ep2737_document

    cfg = load_project_config(CONFIG)
    ctx = {
        "bom_id": "EP 2737",
        "title": "WELDED PLANE PIPE FLANGE",
        "customer": "Test Customer",
        "prepared_by": {"name": "A", "role": "B"},
        "checked_by": {"name": "C", "role": "D"},
        "approved_by": {"name": "E", "role": "F"},
        "harmonic_x": {
            "direction": "X",
            "peak_displacement_mm": 11.38,
            "peak_frequency_hz": 3.91,
            "narrative": {"observations": ["Peak noted."], "conclusions": ["Review."]},
        },
        "shock": {
            "directions": [
                {"direction": "+X", "key": "plus_x", "max_stress_mpa": 255.0, "max_deformation_mm": 0.066, "fos": 1.8}
            ],
            "narrative": {"conclusions": ["Shock OK."]},
        },
    }
    cfg.sections_enabled = ["harmonic_x", "shock"]
    doc = assemble_ep2737_document(ctx, cfg)

    hx = next(s for s in doc.sections if s.key == "harmonic_x")
    assert any(b.kind == "scalar" for b in hx.blocks)
    assert any(b.kind == "pending" and "Figure 40" in (b.caption or "") for b in hx.blocks)
    assert not any(b.kind == "table" for b in hx.blocks)

    shock = next(s for s in doc.sections if s.key == "shock")
    assert any(b.kind == "table" for b in shock.blocks)


def test_block_render_writes_docx(tmp_path):
    from ansys_report.config import load_project_config
    from ansys_report.report.block_assembler import assemble_ep2737_document
    from ansys_report.report.block_render import render_blocks_document
    from docx import Document

    cfg = load_project_config(CONFIG)
    ctx = {
        "bom_id": "EP 2737",
        "title": "WELDED PLANE PIPE FLANGE",
        "customer": "Test",
        "prepared_by": {"name": "A", "role": "B"},
        "checked_by": {"name": "C", "role": "D"},
        "approved_by": {"name": "E", "role": "F"},
        "modal": {"modes": [{"index": 1, "freq_hz": 421.124}], "narrative": {"conclusions": ["OK"]}},
    }
    cfg.sections_enabled = ["modal"]
    doc_model = assemble_ep2737_document(ctx, cfg)
    out = tmp_path / "blocks.docx"
    render_blocks_document(doc_model, ctx, out)

    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    for table in Document(str(out)).tables:
        for row in table.rows:
            text += "\n" + "\n".join(cell.text for cell in row.cells)
    assert "421.124" in text or "421.12" in text
    assert "Table 14" in text
