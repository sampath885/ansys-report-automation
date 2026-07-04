"""Tests for Excel auto-discover design calculation pipeline."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from ansys_report.excel.auto_discover import (
    auto_discover_design_calcs,
    discover_supplemental_workbooks,
    has_calc_data,
)
from ansys_report.excel.reader import read_design_calcs
from ansys_report.excel.udpl_sheet_table import extract_udpl_sheet_table


def _write_udpl_workbook(path: Path) -> None:
    """UDPL valve workbook with 7-column union bolt + Lame-style sheet."""
    wb = Workbook()
    wb.remove(wb.active)

    lame = wb.create_sheet("Lamis theory_2741")
    lame["C1"] = "Stresses at Inner Surface (r = ri)"
    lame.append([1, "Input", "Internal Body Test Pressure", 1, "MPa", "PO Specifications"])
    lame.append([2, "Input", "Minimum Internal Radius (ri)", 40, "mm", "Dwg"])
    lame.append([11, "OUTPUT", "Comparison Result:", "ACCEPTABLE", "", "Comparison of stresses"])

    union = wb.create_sheet("Union Flange Bolt Design_2741")
    union["A1"] = "BOLT LOADS"
    union.append(["Sl. No.", "IN/OP", "Parameter", "Notation/ Formula", "Value", "Unit", "Reference"])
    union.append([1, "Input", "Working Pressure (P)", "P", 1, "MPa", "PO"])
    union.append([2, "Input", "Bolt Material", "ASTM A 276 S32100", 205, "MPA", "STUD SHEET"])
    union.append([14, "Output", "Bolt Load", "MAX (WM1, WM2)", 13273.23, "", ""])

    effort = wb.create_sheet("Effort required calculations")
    effort.append(["Sr. No", "IN/OP", "Parameter", "Value", "Unit", "Referance"])
    effort.append([1, "INPUT", "Working Pressure (Mpa)", 15.7, "MPa", "PO"])
    effort.append([14, "Output", "Total Force", 55963.114, "N", "DWG"])

    bom = wb.create_sheet("BOM")
    bom.append(["Part No.", "Description", "Material"])
    bom.append([1, "BODY", "STEEL"])

    wb.save(path)


def test_extract_udpl_union_bolt_seven_columns(tmp_path):
    path = tmp_path / "calc.xlsx"
    _write_udpl_workbook(path)
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    table = extract_udpl_sheet_table(wb["Union Flange Bolt Design_2741"], sheet_name="Union Flange Bolt Design_2741")
    wb.close()
    assert table is not None
    assert "IN/OP" in table.headers
    assert "Notation/ Formula" in table.headers
    assert len(table.raw_rows) >= 3
    assert table.raw_rows[0][2] == "Working Pressure (P)"
    assert table.raw_rows[0][4] == "1"


def test_extract_udpl_lame_without_header_row(tmp_path):
    path = tmp_path / "calc.xlsx"
    _write_udpl_workbook(path)
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    table = extract_udpl_sheet_table(wb["Lamis theory_2741"], sheet_name="Lamis theory_2741")
    wb.close()
    assert table is not None
    assert table.raw_rows[0][1].upper() == "INPUT"
    assert "ACCEPTABLE" in table.raw_rows[-1][3]


def test_auto_discover_excludes_bom(tmp_path):
    path = tmp_path / "PWS VALVES CALCULATIONS_9999.xlsx"
    _write_udpl_workbook(path)
    result = auto_discover_design_calcs(path)
    sheet_names = {s.sheet_name for s in result.discovered_sections}
    assert "BOM" not in sheet_names
    assert len(result.discovered_sections) >= 3


def test_auto_discover_preserves_raw_columns(tmp_path):
    path = tmp_path / "calc.xlsx"
    _write_udpl_workbook(path)
    result = auto_discover_design_calcs(path)
    union = next(s for s in result.discovered_sections if "Union Flange" in s.sheet_name)
    assert len(union.headers) == 7
    assert len(union.raw_rows[0]) == 7


def test_read_design_calcs_uses_cli_path_not_map(tmp_path):
    primary = tmp_path / "primary.xlsx"
    _write_udpl_workbook(primary)
    map_path = tmp_path / "fake_map.yaml"
    map_path.write_text(
        "variant: ep2737\nworkbooks:\n  design_calcs: missing.xlsx\ntables: {}\n",
        encoding="utf-8",
    )
    result = read_design_calcs(
        primary,
        excel_map_path=map_path,
        case_root=tmp_path,
        excel_mode="auto_then_map",
    )
    assert has_calc_data(result)
    assert result.extraction_source == "auto_discover"


def test_discover_supplemental_bolt_workbook(tmp_path):
    primary = tmp_path / "valve_calcs.xlsx"
    _write_udpl_workbook(primary)
    bolt_wb = Workbook()
    ws = bolt_wb.active
    ws.title = "Bolt preload data"
    ws.append(["Sl. No.", "IN/OP", "Parameter", "Value", "Unit", "Reference"])
    ws.append([1, "Input", "Preload", 5000, "N", "Calc"])
    bolt_path = tmp_path / "Bolt pre load.xlsx"
    bolt_wb.save(bolt_path)
    found = discover_supplemental_workbooks(primary, search_dir=tmp_path)
    assert bolt_path.resolve() in found


def test_assembled_design_calcs_raw_table(tmp_path):
    from ansys_report.config import ProjectConfig, PersonRole
    from ansys_report.report.block_assembler import assemble_ep2737_document

    path = tmp_path / "valve.xlsx"
    _write_udpl_workbook(path)
    cfg = ProjectConfig(
        bom_id="T",
        title="VALVE",
        customer="C",
        prepared_by=PersonRole(name="a", role="b"),
        checked_by=PersonRole(name="a", role="b"),
        approved_by=PersonRole(name="a", role="b"),
        project_dir=tmp_path,
        sections_enabled=["design_calcs"],
        excel_calcs=str(path),
        excel_mode="auto",
    )
    calcs = read_design_calcs(path, excel_mode="auto")
    context = {
        "design_calcs": {
            "discovered_sections": [s.model_dump() for s in calcs.discovered_sections],
            "narrative": {"conclusions": ["OK"], "observations": [], "recommendations": []},
        },
        "methodology": {"fatigue_theory": "Fatigue theory paragraph."},
        "reference_tables": {},
    }
    doc = assemble_ep2737_document(context, cfg)
    design = next(s for s in doc.sections if s.key == "design_calcs")
    tables = [b for b in design.blocks if b.kind == "table"]
    assert len(tables) >= 3
    union_table = next(
        t for t in tables if t.headers and any("Notation" in h for h in t.headers)
    )
    assert len(union_table.headers) == 7
    assert union_table.rows[0][2] == "Working Pressure (P)"
