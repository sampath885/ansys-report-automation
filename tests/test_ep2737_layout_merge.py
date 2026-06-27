"""Tests for EP2737 layout merge preserving embedded figures."""

from __future__ import annotations

import re
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches

from ansys_report.report.ep2737_layout import (
    _patch_content_types,
    _replace_relationship_id,
    merge_reference_front_matter,
)

REPO = Path(__file__).resolve().parents[1]
LAYOUT = REPO / "templates" / "EP2737_layout_template.docx"
STYLE_SHELL = REPO / "templates" / "EP2737_style_shell.docx"


def _embed_ids_and_missing_rels(docx_path: Path) -> tuple[int, int, int]:
    with zipfile.ZipFile(docx_path, "r") as z:
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
        doc = z.read("word/document.xml").decode("utf-8")
    rel_ids = set(re.findall(r'Id="(rId\d+)"', rels))
    embed_ids = re.findall(r'r:embed="(rId\d+)"', doc)
    unique = set(embed_ids)
    return len(embed_ids), len(unique), len(unique - rel_ids)


def _ordered_media_hashes(docx_path: Path) -> list[bytes]:
    with zipfile.ZipFile(docx_path, "r") as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")
        rels_xml = z.read("word/_rels/document.xml.rels").decode("utf-8")
    embed_ids = re.findall(r'r:embed="(rId\d+)"', doc_xml)
    targets = dict(re.findall(r'Id="(rId\d+)"[^>]+Target="([^"]+)"', rels_xml))
    hashes: list[bytes] = []
    with zipfile.ZipFile(docx_path, "r") as z:
        for embed_id in embed_ids:
            target = targets.get(embed_id)
            if not target:
                continue
            part = target if target.startswith("word/") else f"word/{target.lstrip('/')}"
            if part in z.namelist():
                hashes.append(z.read(part))
    return hashes


def _make_minimal_reference(path: Path, *, marker: str = "Revision log") -> None:
    doc = Document()
    doc.add_paragraph("Report cover")
    doc.add_page_break()
    doc.add_paragraph(marker)
    doc.add_paragraph("Discarded reference sample body")
    doc.save(path)


def test_layout_merge_preserves_generated_image_relationships(tiny_png):
    assert LAYOUT.exists()

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        generated = tmp_path / "generated.docx"
        output = tmp_path / "merged.docx"

        doc = Document()
        doc.add_heading("Revision log", level=1)
        doc.add_paragraph("Generated body")
        doc.add_picture(str(tiny_png), width=Inches(2))
        doc.save(generated)

        merge_reference_front_matter(LAYOUT, generated, output)

        embeds, _, missing = _embed_ids_and_missing_rels(output)
        assert embeds >= 1
        assert missing == 0


def test_layout_merge_preserves_multiple_generated_images(tiny_png):
    assert LAYOUT.exists()

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        generated = tmp_path / "generated.docx"
        output = tmp_path / "merged.docx"

        doc = Document()
        doc.add_heading("Revision log", level=1)
        for idx in range(5):
            doc.add_paragraph(f"Figure block {idx}")
            doc.add_picture(str(tiny_png), width=Inches(2))
        doc.save(generated)

        merge_reference_front_matter(LAYOUT, generated, output)

        embeds, _, missing = _embed_ids_and_missing_rels(output)
        assert embeds >= 5
        assert missing == 0
        with zipfile.ZipFile(output) as zout:
            doc = zout.read("word/document.xml").decode("utf-8")
        assert doc.count("pic:pic") >= 5


def test_replace_relationship_id_does_not_corrupt_longer_ids():
    xml = (
        '<p r:embed="rId1"/> middle '
        '<p r:embed="rId10"/> '
        '<p r:embed="rId11"/> '
        '<p r:link="rId2"/> '
        '<p r:link="rId20"/>'
    )
    out = _replace_relationship_id(xml, "rId1", "rId500")
    assert 'r:embed="rId500"' in out
    assert 'r:embed="rId10"' in out
    assert 'r:embed="rId11"' in out
    assert 'r:embed="rId1"' not in out

    out2 = _replace_relationship_id(out, "rId2", "rId501")
    assert 'r:link="rId501"' in out2
    assert 'r:link="rId20"' in out2
    assert 'r:link="rId2"' not in out2


def test_patch_content_types_adds_generated_media():
    base = b"""<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
</Types>"""
    patched = _patch_content_types(base, ["word/media/gen_0001.png"])
    text = patched.decode("utf-8")
    assert '/word/media/gen_0001.png"' in text
    assert "image/png" in text


def test_layout_merge_preserves_fifteen_distinct_images(tmp_path, tiny_png):
    """Regression: ascending str.replace turned rId10 into rId500 after remapping rId1."""
    assert LAYOUT.exists()

    pngs: list[Path] = []
    for idx in range(15):
        path = tmp_path / f"fig_{idx:02d}.png"
        path.write_bytes(tiny_png.read_bytes() + bytes([idx]))
        pngs.append(path)

    generated = tmp_path / "generated.docx"
    output = tmp_path / "merged.docx"

    doc = Document()
    doc.add_heading("Revision log", level=1)
    for idx, png in enumerate(pngs):
        doc.add_paragraph(f"Figure block {idx}")
        doc.add_picture(str(png), width=Inches(2))
    doc.save(generated)

    expected = _ordered_media_hashes(generated)
    assert len(expected) >= 15

    merge_reference_front_matter(LAYOUT, generated, output)

    embeds, _, missing = _embed_ids_and_missing_rels(output)
    assert embeds >= 15
    assert missing == 0

    with zipfile.ZipFile(output, "r") as zout:
        doc_xml = zout.read("word/document.xml").decode("utf-8")
        rels_xml = zout.read("word/_rels/document.xml.rels").decode("utf-8")
        ct_xml = zout.read("[Content_Types].xml").decode("utf-8")
        gen_media = [n for n in zout.namelist() if "word/media/gen_" in n]

    assert "rId500" not in doc_xml
    assert "rId505" not in doc_xml
    assert len(gen_media) >= 15
    assert "/word/media/gen_" in ct_xml

    # Body embed targets must use gen_* parts, not stale imageN.png reuse.
    body_embed_ids = re.findall(r'r:embed="(rId\d+)"', doc_xml)
    targets = dict(re.findall(r'Id="(rId\d+)"[^>]+Target="([^"]+)"', rels_xml))
    body_targets = [targets[rid] for rid in body_embed_ids if rid in targets]
    assert body_targets
    assert all("gen_" in t for t in body_targets[-len(expected) :])

    merged_hashes = _ordered_media_hashes(output)
    body_start = len(merged_hashes) - len(expected)
    assert body_start >= 0
    assert merged_hashes[body_start:] == expected


def test_merge_uses_generated_bytes_not_stale_reference_media(tmp_path, tiny_png):
    """When reference and generated share imageN.png names, body must use generated bytes."""
    reference = tmp_path / "reference.docx"
    generated = tmp_path / "generated.docx"
    output = tmp_path / "merged.docx"

    stale = tmp_path / "stale.png"
    fresh = tmp_path / "fresh.png"
    stale.write_bytes(tiny_png.read_bytes() + b"STALE")
    fresh.write_bytes(tiny_png.read_bytes() + b"FRESH")

    ref_doc = Document()
    ref_doc.add_paragraph("Cover")
    ref_doc.add_page_break()
    ref_doc.add_paragraph("Revision log")
    ref_doc.add_paragraph("Old reference figure")
    ref_doc.add_picture(str(stale), width=Inches(2))
    ref_doc.save(reference)

    gen_doc = Document()
    gen_doc.add_heading("Revision log", level=1)
    gen_doc.add_paragraph("Figure 6: Mesh")
    gen_doc.add_picture(str(fresh), width=Inches(2))
    gen_doc.save(generated)

    merge_reference_front_matter(reference, generated, output)

    fresh_hash = fresh.read_bytes()
    merged_hashes = _ordered_media_hashes(output)
    assert fresh_hash in merged_hashes
    assert stale.read_bytes() not in merged_hashes[-1:]


def test_layout_merge_with_style_shell_generated_doc(tiny_png, tmp_path):
    """Production path: generated from EP2737 style shell, merged into layout template."""
    if not LAYOUT.exists() or not STYLE_SHELL.exists():
        return

    from ansys_report.report.ep2737_styles import open_document_from_shell

    pngs: list[Path] = []
    for idx in range(8):
        path = tmp_path / f"shell_fig_{idx:02d}.png"
        path.write_bytes(tiny_png.read_bytes() + bytes([idx + 10]))
        pngs.append(path)

    generated = tmp_path / "generated.docx"
    output = tmp_path / "merged.docx"

    doc = open_document_from_shell(STYLE_SHELL)
    doc.add_heading("Revision log", level=1)
    for idx, png in enumerate(pngs):
        doc.add_paragraph(f"Shell figure {idx}")
        doc.add_picture(str(png), width=Inches(2))
    doc.save(generated)

    expected = _ordered_media_hashes(generated)
    assert len(expected) >= 8

    merge_reference_front_matter(LAYOUT, generated, output)

    embeds, _, missing = _embed_ids_and_missing_rels(output)
    assert embeds >= 8
    assert missing == 0

    merged_hashes = _ordered_media_hashes(output)
    body_start = len(merged_hashes) - len(expected)
    assert merged_hashes[body_start:] == expected
