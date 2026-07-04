"""Read design calculation tables from Excel workbooks."""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import load_workbook

from ansys_report.models import CalcRow, DesignCalcsResult

logger = logging.getLogger(__name__)

TABLE_SHEETS = {
    "bolt_load": "Bolt Load",
    "flange_moments": "Flange Moments",
    "effort": "Effort",
    "end_flange": "End Flange",
}

HEADER_ROW = 1


def read_design_calcs(
    excel_path: Path,
    *,
    excel_map_path: Path | None = None,
    case_root: Path | None = None,
    excel_bolt_preload: Path | None = None,
    export_image_dir: Path | None = None,
    fos_target: float = 1.5,
    excel_mode: str = "auto_then_map",
) -> DesignCalcsResult:
    """Load design calculations: auto-discover first, then optional EP2737 map fallback."""
    excel_path = Path(excel_path)

    if excel_mode in ("auto", "auto_then_map") and excel_path.exists():
        from ansys_report.excel.auto_discover import (
            auto_discover_design_calcs,
            discover_supplemental_workbooks,
            has_calc_data,
        )

        supplements = discover_supplemental_workbooks(
            excel_path,
            explicit_bolt_preload=excel_bolt_preload,
            search_dir=case_root or excel_path.parent,
        )
        auto_result = auto_discover_design_calcs(
            excel_path,
            supplemental_workbooks=supplements,
            export_image_dir=export_image_dir,
            fos_target=fos_target,
        )
        if has_calc_data(auto_result):
            return auto_result
        if excel_mode == "auto":
            logger.warning(
                "Auto-discover found no calculation tables in %s; trying map/generic fallback",
                excel_path.name,
            )

    if excel_map_path and excel_map_path.exists() and case_root:
        from ansys_report.excel.auto_discover import has_calc_data
        from ansys_report.excel.ep2737 import load_excel_map, read_ep2737_design_calcs

        try:
            load_excel_map(excel_map_path)
            mapped = read_ep2737_design_calcs(case_root, excel_map_path, fos_target=fos_target)
            mapped.extraction_source = mapped.extraction_source or "excel_map"
            if has_calc_data(mapped):
                return mapped
        except Exception as exc:
            logger.warning("EP2737 excel adapter failed (%s); falling back to generic reader", exc)

    if not excel_path.exists():
        logger.warning("Excel file not found: %s", excel_path)
        return DesignCalcsResult(extraction_source="missing")

    wb = load_workbook(excel_path, data_only=True)
    result = DesignCalcsResult(extraction_source="generic")

    for attr, sheet_name in TABLE_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            continue
        rows = _read_table_sheet(wb[sheet_name])
        setattr(result, attr, rows)

    _read_named_ranges(wb, result)
    wb.close()
    return result


def _read_table_sheet(ws) -> list[CalcRow]:
    rows: list[CalcRow] = []
    headers = [cell.value for cell in ws[HEADER_ROW]]
    col_map = _map_columns(headers)
    if not col_map:
        return rows

    for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        if not row or row[0] is None:
            continue
        rows.append(
            CalcRow(
                label=str(row[col_map.get("label", 0)] or ""),
                symbol=_str(row, col_map, "symbol"),
                formula=_str(row, col_map, "formula"),
                value=row[col_map.get("value", 1)] if "value" in col_map else row[1],
                unit=_str(row, col_map, "unit"),
                source_ref=_str(row, col_map, "source_ref"),
                verdict=_str(row, col_map, "verdict"),
            )
        )
    return rows


def _map_columns(headers: list) -> dict[str, int]:
    mapping: dict[str, int] = {}
    aliases = {
        "label": ("label", "parameter", "description"),
        "symbol": ("symbol",),
        "formula": ("formula",),
        "value": ("value", "result"),
        "unit": ("unit", "units"),
        "source_ref": ("source", "source_ref", "reference"),
        "verdict": ("verdict", "status"),
    }
    for i, h in enumerate(headers):
        if h is None:
            continue
        key = str(h).strip().lower()
        for field, names in aliases.items():
            if key in names and field not in mapping:
                mapping[field] = i
    return mapping


def _str(row, col_map: dict[str, int], field: str) -> str | None:
    if field not in col_map:
        return None
    val = row[col_map[field]]
    return str(val) if val is not None else None


def _read_named_ranges(wb, result: DesignCalcsResult) -> None:
    for name, attr in (
        ("BoltLoad", "bolt_load"),
        ("FlangeMoments", "flange_moments"),
        ("Effort", "effort"),
        ("EndFlange", "end_flange"),
    ):
        if name not in wb.defined_names or getattr(result, attr):
            continue
        dest = wb.defined_names[name].destinations
        for sheet_name, ref in dest:
            ws = wb[sheet_name]
            rows = _read_range(ws, ref)
            if rows:
                setattr(result, attr, rows)


def _read_range(ws, ref: str) -> list[CalcRow]:
    cells = ws[ref]
    if not hasattr(cells, "__iter__"):
        cells = [[cells]]
    rows: list[CalcRow] = []
    for row in cells:
        vals = [c.value for c in row]
        if not vals or vals[0] is None:
            continue
        rows.append(
            CalcRow(
                label=str(vals[0]),
                symbol=str(vals[1]) if len(vals) > 1 and vals[1] else None,
                formula=str(vals[2]) if len(vals) > 2 and vals[2] else None,
                value=vals[3] if len(vals) > 3 else vals[1],
                unit=str(vals[4]) if len(vals) > 4 and vals[4] else None,
                source_ref=str(vals[5]) if len(vals) > 5 and vals[5] else None,
                verdict=str(vals[6]) if len(vals) > 6 and vals[6] else None,
            )
        )
    return rows
