"""Auto-discover UDPL design calculation sheets and paste raw Excel tables into the report."""

from __future__ import annotations

import logging
import re
from pathlib import Path

from openpyxl import load_workbook

from ansys_report.excel.udpl_sheet_table import (
    CATEGORY_ORDER,
    extract_udpl_sheet_table,
    should_exclude_sheet,
    sort_tables,
)
from ansys_report.models import CalcRow, DesignCalcsResult, DiscoveredCalcSection

logger = logging.getLogger(__name__)

LEGACY_CATEGORY_MAP = {
    "union_bolt": "bolt_load",
    "bolt_preload": "bolt_load",
    "bolt_thread": "bolt_load",
    "flange_rigidity": "flange_moments",
    "wall_thickness": "end_flange",
    "lame": "end_flange",
    "pressure_class": "end_flange",
    "effort": "effort",
    "spindle": "effort",
}


def auto_discover_design_calcs(
    primary_workbook: Path,
    *,
    supplemental_workbooks: list[Path] | None = None,
    export_image_dir: Path | None = None,
    fos_target: float = 1.5,
    min_rows: int = 2,
) -> DesignCalcsResult:
    """Scan workbook(s) and extract calculation tables verbatim from Excel."""
    del fos_target  # UDPL paste does not recompute verdicts.
    if not primary_workbook.exists():
        logger.warning("Primary design calcs workbook not found: %s", primary_workbook)
        return DesignCalcsResult(extraction_source="missing")

    workbooks = [primary_workbook.resolve()]
    for path in supplemental_workbooks or []:
        resolved = path.resolve()
        if resolved.exists() and resolved not in workbooks:
            workbooks.append(resolved)

    sections: list[DiscoveredCalcSection] = []
    legacy: dict[str, list[CalcRow]] = {
        "end_flange": [],
        "flange_moments": [],
        "bolt_load": [],
        "effort": [],
    }

    for workbook_path in workbooks:
        wb = load_workbook(workbook_path, data_only=True)
        try:
            sheet_names = list(wb.sheetnames)
            for sheet_name in sheet_names:
                if should_exclude_sheet(sheet_name, sheet_names):
                    continue
                ws = wb[sheet_name]
                table = extract_udpl_sheet_table(ws, sheet_name=sheet_name)
                if table is not None and len(table.raw_rows) >= min_rows:
                    section = DiscoveredCalcSection(
                        key=_slug(table.category, sheet_name),
                        title=table.title,
                        sheet_name=table.sheet_name,
                        workbook=workbook_path.name,
                        category=table.category,
                        headers=table.headers,
                        raw_rows=table.raw_rows,
                        subtitle=table.subtitle,
                    )
                    sections.append(section)
                    bucket = LEGACY_CATEGORY_MAP.get(table.category)
                    if bucket:
                        legacy[bucket].extend(_legacy_rows_from_raw(table))
                    continue

                if export_image_dir is not None:
                    image_path = _try_export_sheet_png(
                        workbook_path, sheet_name, export_image_dir
                    )
                    if image_path:
                        sections.append(
                            DiscoveredCalcSection(
                                key=_slug("general", sheet_name),
                                title=_title_from_sheet(sheet_name),
                                sheet_name=sheet_name,
                                workbook=workbook_path.name,
                                category="general",
                                raw_rows=[],
                                headers=[],
                                image_path=str(image_path),
                            )
                        )
        finally:
            wb.close()

    sections = _sort_discovered_sections(sections)
    if not sections and not any(legacy.values()):
        return DesignCalcsResult(extraction_source="auto_discover_empty")

    return DesignCalcsResult(
        end_flange=legacy["end_flange"],
        flange_moments=legacy["flange_moments"],
        bolt_load=legacy["bolt_load"],
        effort=legacy["effort"],
        discovered_sections=sections,
        extraction_source="auto_discover",
    )


def discover_supplemental_workbooks(
    primary_workbook: Path,
    *,
    explicit_bolt_preload: Path | None = None,
    search_dir: Path | None = None,
) -> list[Path]:
    """Find bolt-preload workbooks in the same folder without user configuration."""
    found: list[Path] = []
    seen: set[Path] = set()

    def _add(path: Path | None) -> None:
        if path is None or not path.exists():
            return
        resolved = path.resolve()
        if resolved in seen:
            return
        seen.add(resolved)
        found.append(resolved)

    _add(explicit_bolt_preload)
    folder = search_dir or primary_workbook.parent
    if not folder.exists():
        return found

    for candidate in sorted(folder.glob("*.xlsx")):
        if candidate.resolve() == primary_workbook.resolve():
            continue
        name = candidate.name.lower()
        if any(token in name for token in ("bolt", "preload", "pretension", "pre load", "pre-load")):
            _add(candidate)
    return found


def has_calc_data(result: DesignCalcsResult) -> bool:
    if result.discovered_sections:
        return True
    return any([result.bolt_load, result.flange_moments, result.effort, result.end_flange])


def _legacy_rows_from_raw(table) -> list[CalcRow]:
    """Best-effort CalcRow list for legacy fatigue summary (label/value only)."""
    headers = [h.lower() for h in table.headers]
    label_idx = _col_index(headers, ("parameter", "particulars"))
    value_idx = _col_index(headers, ("value", "result"))
    unit_idx = _col_index(headers, ("unit", "units"))
    ref_idx = _col_index(headers, ("reference", "referance", "source"))
    if label_idx is None or value_idx is None:
        return []

    rows: list[CalcRow] = []
    for raw in table.raw_rows:
        if label_idx >= len(raw) or value_idx >= len(raw):
            continue
        label = raw[label_idx].strip()
        value_text = raw[value_idx].strip()
        if not label or not value_text:
            continue
        value: float | str = value_text
        try:
            value = float(value_text.replace(",", ""))
        except ValueError:
            pass
        rows.append(
            CalcRow(
                label=label,
                value=value,
                unit=raw[unit_idx].strip() if unit_idx is not None and unit_idx < len(raw) else None,
                source_ref=raw[ref_idx].strip() if ref_idx is not None and ref_idx < len(raw) else None,
            )
        )
    return rows


def _col_index(headers: list[str], names: tuple[str, ...]) -> int | None:
    for idx, header in enumerate(headers):
        if any(name in header for name in names):
            return idx
    return None


def _title_from_sheet(sheet_name: str) -> str:
    text = re.sub(r"_2741$|_cn8$", "", sheet_name, flags=re.I)
    text = re.sub(r"[_\-]+", " ", text.strip())
    return re.sub(r"\s+", " ", text).strip()


def _slug(category: str, sheet_name: str) -> str:
    base = re.sub(r"[^a-z0-9]+", "_", sheet_name.lower()).strip("_")
    return f"{category}_{base}" if base else category


def _sort_discovered_sections(sections: list[DiscoveredCalcSection]) -> list[DiscoveredCalcSection]:
    order = {name: idx for idx, name in enumerate(CATEGORY_ORDER)}

    def _key(section: DiscoveredCalcSection) -> tuple[int, str]:
        return (order.get(section.category, len(CATEGORY_ORDER)), section.title.lower())

    return sorted(sections, key=_key)


def _try_export_sheet_png(
    workbook_path: Path,
    sheet_name: str,
    output_dir: Path,
) -> Path | None:
    output_dir.mkdir(parents=True, exist_ok=True)
    slug = re.sub(r"[^a-z0-9]+", "_", sheet_name.lower()).strip("_")
    dest = output_dir / f"{workbook_path.stem}_{slug}.png"
    if dest.exists():
        return dest

    try:
        import win32com.client  # type: ignore[import-untyped]
    except ImportError:
        logger.debug("win32com unavailable; skipping sheet PNG export for %s", sheet_name)
        return None

    excel = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Open(str(workbook_path.resolve()))
        try:
            ws = wb.Worksheets(sheet_name)
            ws.Activate()
            used = ws.UsedRange
            used.CopyPicture(Format=2)
            chart_obj = ws.ChartObjects().Add(0, 0, used.Width, used.Height)
            chart = chart_obj.Chart
            chart.Paste()
            chart.Export(str(dest.resolve()))
            chart_obj.Delete()
        finally:
            wb.Close(SaveChanges=False)
    except Exception as exc:
        logger.warning("Sheet PNG export failed for %r: %s", sheet_name, exc)
        if dest.exists():
            dest.unlink(missing_ok=True)
        return None
    finally:
        if excel is not None:
            excel.Quit()

    return dest if dest.exists() else None
