"""EP2737-specific Excel layout adapter."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

from ansys_report.models import CalcRow, DesignCalcsResult

logger = logging.getLogger(__name__)


class ExcelMapError(ValueError):
    pass


def load_excel_map(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}
    if data.get("variant") != "ep2737":
        raise ExcelMapError(f"Unsupported excel map variant: {data.get('variant')}")
    return data


def read_ep2737_design_calcs(
    case_root: Path,
    map_path: Path,
    *,
    fos_target: float = 1.5,
) -> DesignCalcsResult:
    """Read EP2737 design calculation tables using ``ep2737_excel_map.yaml``."""
    cfg = load_excel_map(map_path)
    workbooks = _resolve_workbooks(case_root, cfg.get("workbooks", {}))

    result = DesignCalcsResult()
    for table_key, table_cfg in cfg.get("tables", {}).items():
        wb_key = table_cfg["workbook"]
        wb_path = workbooks.get(wb_key)
        if wb_path is None or not wb_path.exists():
            logger.warning("Workbook missing for %s: %s", table_key, wb_key)
            continue
        rows = _read_table(wb_path, table_cfg, fos_target=fos_target)
        if hasattr(result, table_key):
            setattr(result, table_key, rows)
        else:
            logger.warning("Unknown DesignCalcsResult field: %s", table_key)
    return result


def _resolve_workbooks(case_root: Path, spec: dict[str, str]) -> dict[str, Path]:
    resolved: dict[str, Path] = {}
    for key, rel in spec.items():
        path = Path(rel)
        resolved[key] = path if path.is_absolute() else case_root / rel
    return resolved


def _read_table(workbook_path: Path, table_cfg: dict[str, Any], *, fos_target: float) -> list[CalcRow]:
    wb = load_workbook(workbook_path, data_only=True)
    sheet_name = table_cfg["sheet"]
    if sheet_name not in wb.sheetnames:
        logger.warning("Sheet %r not in %s", sheet_name, workbook_path.name)
        wb.close()
        return []

    ws = wb[sheet_name]
    layout = table_cfg["layout"]
    if layout == "parameters_bcde":
        rows = _read_parameters_bcde(ws, table_cfg)
    elif layout == "label_value_unit":
        rows = _read_label_value_unit(ws, table_cfg)
    elif layout == "parameter_column":
        if table_cfg.get("report_table", "").startswith("Table 27"):
            rows = _read_fatigue_effort_table(ws, table_cfg)
        else:
            rows = _read_parameter_column(ws, table_cfg)
    else:
        raise ExcelMapError(f"Unknown layout: {layout}")

    wb.close()
    return _apply_verdicts(rows, fos_target=fos_target)


def _cell(row: tuple, col: int) -> Any:
    if col >= len(row):
        return None
    return row[col]


def _clean_label(text: Any) -> str | None:
    if text is None:
        return None
    label = str(text).strip()
    return label or None


def _read_parameters_bcde(ws, cfg: dict[str, Any]) -> list[CalcRow]:
    header_row = cfg.get("header_row", 3)
    label_col = cfg["label_col"]
    unit_col = cfg["unit_col"]
    value_col = cfg["value_col"]
    source_col = cfg.get("source_col")
    filters = [s.upper() for s in cfg.get("label_contains", [])]

    rows: list[CalcRow] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i <= header_row:
            continue
        label = _clean_label(_cell(row, label_col))
        if not label:
            continue
        if filters and not any(f in label.upper() for f in filters):
            continue
        value = _cell(row, value_col)
        if value is None and not filters:
            continue
        rows.append(
            CalcRow(
                label=label,
                value=value if value is not None else "",
                unit=_clean_label(_cell(row, unit_col)),
                source_ref=_clean_label(_cell(row, source_col)) if source_col is not None else None,
                verdict=_infer_verdict(label, value),
            )
        )
    return rows


def _read_label_value_unit(ws, cfg: dict[str, Any]) -> list[CalcRow]:
    label_col = cfg["label_col"]
    value_col = cfg["value_col"]
    unit_col = cfg.get("unit_col")
    per_bolt_col = cfg.get("per_bolt_col")
    source_col = cfg.get("source_col")

    rows: list[CalcRow] = []
    for row in ws.iter_rows(values_only=True):
        label = _clean_label(_cell(row, label_col))
        if not label:
            continue
        value = _cell(row, value_col)
        if value is None:
            continue
        per_bolt = _cell(row, per_bolt_col) if per_bolt_col is not None else None
        formula = f"per bolt = {per_bolt}" if per_bolt is not None else None
        rows.append(
            CalcRow(
                label=label,
                value=value,
                unit=_clean_label(_cell(row, unit_col)) if unit_col is not None else None,
                formula=formula,
                source_ref=_clean_label(_cell(row, source_col)) if source_col is not None else None,
            )
        )
    return rows


def _read_fatigue_effort_table(ws, cfg: dict[str, Any]) -> list[CalcRow]:
    """Read EP2737 fatigue/bolt effort sheet (without Gasket layout)."""
    header_row = cfg.get("header_row", 5)
    label_col = cfg["label_col"]
    value_col = cfg["value_col"]
    source_col = cfg.get("source_col")
    data_start = cfg.get("data_start_row", header_row + 1)

    rows: list[CalcRow] = []
    bolt_symbol: str | None = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == header_row:
            bolt_symbol = _clean_label(_cell(row, value_col))
            continue
        if i < data_start:
            continue
        label = _clean_label(_cell(row, label_col))
        value = _cell(row, value_col)
        source_ref = _clean_label(_cell(row, source_col)) if source_col is not None else None
        if label is None and value is None:
            continue
        if label is None and value is not None:
            label = ""
        if value is None:
            continue
        rows.append(
            CalcRow(
                label=label,
                symbol=bolt_symbol,
                value=value,
                unit=_infer_fatigue_unit(label),
                source_ref=source_ref,
            )
        )
    return rows


def _infer_fatigue_unit(label: str) -> str | None:
    lower = label.lower()
    if "diameter" in lower or "pitch" in lower:
        return "mm"
    if lower.startswith("p%"):
        return None
    if "preload" in lower:
        return "N"
    if "σy" in label or "yield" in lower:
        return "MPa"
    if label == "":
        return "mm²"
    return None


def _read_parameter_column(ws, cfg: dict[str, Any]) -> list[CalcRow]:
    header_row = cfg.get("header_row", 5)
    data_start = cfg.get("data_start_row", header_row + 1)
    label_col = cfg["label_col"]
    value_col = cfg["value_col"]
    source_col = cfg.get("source_col")

    rows: list[CalcRow] = []
    header = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == header_row:
            header = _clean_label(_cell(row, value_col))
        if i < data_start:
            continue
        label = _clean_label(_cell(row, label_col))
        value = _cell(row, value_col)
        if not label and value is None:
            continue
        if not label and value is not None:
            label = header or "Value"
        if label and value is not None:
            rows.append(
                CalcRow(
                    label=label,
                    symbol=header,
                    value=value,
                    source_ref=_clean_label(_cell(row, source_col)) if source_col is not None else None,
                )
            )
    return rows


def _infer_verdict(label: str, value: Any) -> str | None:
    upper = label.upper()
    if "ACCEPTED" in str(value).upper():
        return "ACCEPTABLE"
    if value == "ACCEPTED" or (isinstance(value, str) and value.strip().upper() == "ACCEPTED"):
        return "ACCEPTABLE"
    if "ST<" in upper or "TANGENTIAL FLANGE STRESS" in upper:
        if isinstance(value, str) and "ACCEPT" in value.upper():
            return "ACCEPTABLE"
    return None


def _apply_verdicts(rows: list[CalcRow], *, fos_target: float) -> list[CalcRow]:
    updated: list[CalcRow] = []
    for row in rows:
        if row.verdict:
            updated.append(row)
            continue
        if "FACTOR OF SAFETY" in row.label.upper() or row.label.strip().upper() == "FOS":
            try:
                fos = float(row.value)
                verdict = "ACCEPTABLE" if fos >= fos_target else "NOT ACCEPTABLE"
                updated.append(row.model_copy(update={"verdict": verdict}))
                continue
            except (TypeError, ValueError):
                pass
        updated.append(row)
    return updated


def summarize_for_golden(result: DesignCalcsResult) -> dict[str, Any]:
    """Key scalars for golden-file comparison."""

    def _find(rows: list[CalcRow], *needles: str) -> CalcRow | None:
        for row in rows:
            upper = row.label.upper()
            if any(n.upper() in upper for n in needles):
                return row
        return None

    fos = _find(result.end_flange, "FACTOR OF SAFETY", "FOS")
    tr = _find(result.end_flange, "REQUIRED FLANGE THICKNESS")
    st = _find(result.end_flange, "TANGENTIAL FLANGE STRESS, ST")
    stress_verdict = _find(result.end_flange, "tangential flange stress (st)")
    verdict_value = None
    if stress_verdict:
        if isinstance(stress_verdict.value, str):
            verdict_value = stress_verdict.value
        elif stress_verdict.verdict:
            verdict_value = stress_verdict.verdict

    bolts = _find(result.bolt_load, "NO  BOLTS", "NO OF BOLTS")
    pressure = _find(result.bolt_load, "WORKING PRESSURE")
    preload = _find(result.effort, "PRELOAD")
    bolt_size = next((r.symbol for r in result.effort if r.symbol), None)
    yield_row = _find(result.effort, "σy", "yield", "S321")

    return {
        "phase": "4_excel",
        "end_flange": {
            "fos": float(fos.value) if fos and fos.value is not None else None,
            "required_thickness_mm": float(tr.value) if tr and tr.value is not None else None,
            "tangential_stress_mpa": float(st.value) if st and st.value is not None else None,
            "stress_verdict": verdict_value,
        },
        "bolt_load": {
            "bolt_count": int(bolts.value) if bolts and bolts.value is not None else None,
            "design_pressure_mpa": float(pressure.value) if pressure and pressure.value is not None else None,
        },
        "effort": {
            "preload_n": float(preload.value) if preload and preload.value is not None else None,
            "bolt_size": bolt_size,
            "yield_mpa": float(yield_row.value) if yield_row and yield_row.value is not None else None,
        },
        "row_counts": {
            "end_flange": len(result.end_flange),
            "flange_moments": len(result.flange_moments),
            "bolt_load": len(result.bolt_load),
            "effort": len(result.effort),
        },
    }
