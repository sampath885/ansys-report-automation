"""EP1581 Word formatting aligned with EP1581_29APR_R0 reference DOCX."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips
from docx.table import Table

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# Portrait A4 (twips) — used for every page in generated EP1581 reports.
EP1581_PORTRAIT = (11906, 16838)
EP1581_LANDSCAPE = (16839, 11907)
EP1581_MARGINS = {
    "top": 1417,
    "bottom": 1417,
    "left": 1701,
    "right": 1701,
    "header": 520,
    "footer": 300,
    "gutter": 0,
}

EP1581_BODY_FONT = "Times New Roman"
EP1581_BODY_SIZE_PT = 11


def add_body_paragraph(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="Normal")
    run = paragraph.add_run(text)
    _apply_body_font(run)
    return paragraph


def add_section_break(doc: Document, orientation: str = "landscape") -> None:
    """Insert a section break matching reference portrait/landscape pages."""
    width, height = EP1581_LANDSCAPE if orientation.lower() == "landscape" else EP1581_PORTRAIT
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    if orientation.lower() == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
    else:
        section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Twips(width)
    section.page_height = Twips(height)
    section.top_margin = Twips(EP1581_MARGINS["top"])
    section.bottom_margin = Twips(EP1581_MARGINS["bottom"])
    section.left_margin = Twips(EP1581_MARGINS["left"])
    section.right_margin = Twips(EP1581_MARGINS["right"])
    section.header_distance = Twips(EP1581_MARGINS["header"])
    section.footer_distance = Twips(EP1581_MARGINS["footer"])
    section.gutter = Twips(EP1581_MARGINS["gutter"])


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> Table | None:
    """Full-width Table Grid with bold header row (reference layout)."""
    if not headers:
        return None
    if not rows:
        add_body_paragraph(doc, "(no data)")
        return None

    col_count = max(len(headers), max((len(r) for r in rows), default=0))
    norm_headers = headers + [""] * (col_count - len(headers))
    table = doc.add_table(rows=1 + len(rows), cols=col_count)
    table.style = "Table Grid"
    _set_table_full_width(table)
    _set_equal_column_widths(table, col_count)

    for col_idx, header in enumerate(norm_headers):
        cell = table.rows[0].cells[col_idx]
        cell.text = header
        _bold_cell(cell)
        _apply_body_font_to_cell(cell)

    for row_idx, row in enumerate(rows, start=1):
        padded = row + [""] * (col_count - len(row))
        for col_idx, value in enumerate(padded[:col_count]):
            cell = table.rows[row_idx].cells[col_idx]
            cell.text = str(value)
            _apply_body_font_to_cell(cell)

    return table


def _apply_body_font(run) -> None:
    run.font.name = EP1581_BODY_FONT
    run.font.size = Pt(EP1581_BODY_SIZE_PT)


def _apply_body_font_to_cell(cell) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            _apply_body_font(run)


def _bold_cell(cell) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True


def _set_table_full_width(table: Table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")

    tbl_look = tbl_pr.find(qn("w:tblLook"))
    if tbl_look is None:
        tbl_look = OxmlElement("w:tblLook")
        tbl_pr.append(tbl_look)
    tbl_look.set(qn("w:firstRow"), "1")
    tbl_look.set(qn("w:lastRow"), "0")
    tbl_look.set(qn("w:firstColumn"), "1")
    tbl_look.set(qn("w:lastColumn"), "0")
    tbl_look.set(qn("w:noHBand"), "0")
    tbl_look.set(qn("w:noVBand"), "1")
    tbl_look.set(qn("w:val"), "04A0")


def _set_equal_column_widths(table: Table, col_count: int) -> None:
    if col_count <= 0:
        return
    pct = str(5000 // col_count)
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), pct)
            tc_w.set(qn("w:type"), "pct")


def add_caption(doc: Document, text: str) -> None:
    cap = doc.add_paragraph(text, style="Caption")
    cap.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER


def add_centered_picture(doc: Document, image_path: str, width_in: float) -> None:
    """Embed a figure centered on the page (EP1581 layout)."""
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = paragraph.add_run()
    run.add_picture(image_path, width=Inches(width_in))


def _force_portrait_a4_on_element(pg_sz: ET.Element) -> None:
    width, height = EP1581_PORTRAIT
    pg_sz.set(f"{{{W_NS}}}w", str(width))
    pg_sz.set(f"{{{W_NS}}}h", str(height))
    orient_key = f"{{{W_NS}}}orient"
    if orient_key in pg_sz.attrib:
        del pg_sz.attrib[orient_key]


def _apply_ep1581_margins(sect_pr: ET.Element) -> None:
    pg_mar = sect_pr.find(f"{{{W_NS}}}pgMar")
    if pg_mar is None:
        pg_mar = ET.SubElement(sect_pr, f"{{{W_NS}}}pgMar")
    pg_mar.set(f"{{{W_NS}}}top", str(EP1581_MARGINS["top"]))
    pg_mar.set(f"{{{W_NS}}}bottom", str(EP1581_MARGINS["bottom"]))
    pg_mar.set(f"{{{W_NS}}}left", str(EP1581_MARGINS["left"]))
    pg_mar.set(f"{{{W_NS}}}right", str(EP1581_MARGINS["right"]))
    pg_mar.set(f"{{{W_NS}}}header", str(EP1581_MARGINS["header"]))
    pg_mar.set(f"{{{W_NS}}}footer", str(EP1581_MARGINS["footer"]))
    pg_mar.set(f"{{{W_NS}}}gutter", str(EP1581_MARGINS["gutter"]))


def normalize_document_xml_portrait_a4(document_xml: str) -> str:
    """Force every section in document.xml to portrait A4 with EP1581 margins."""
    root = ET.fromstring(document_xml)
    for sect_pr in root.iter(f"{{{W_NS}}}sectPr"):
        pg_sz = sect_pr.find(f"{{{W_NS}}}pgSz")
        if pg_sz is None:
            pg_sz = ET.SubElement(sect_pr, f"{{{W_NS}}}pgSz")
        _force_portrait_a4_on_element(pg_sz)
        _apply_ep1581_margins(sect_pr)
    ET.register_namespace("w", W_NS)
    return '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + ET.tostring(
        root, encoding="unicode", xml_declaration=False
    )


def normalize_docx_portrait_a4(path: Path) -> None:
    """Rewrite all sections in a DOCX to portrait A4 (fixes shrunk/landscape merged pages)."""
    source = path.read_bytes()
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source), "r") as zin:
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename == "word/document.xml":
                    data = normalize_document_xml_portrait_a4(
                        data.decode("utf-8")
                    ).encode("utf-8")
                zout.writestr(info, data)
    path.write_bytes(buffer.getvalue())
