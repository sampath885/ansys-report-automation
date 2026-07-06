"""Render assembled content blocks to a Word document."""

from __future__ import annotations

import logging
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt

from ansys_report.report.blocks import RenderBlock, RenderDocument, RenderSection
from ansys_report.report.docx_sanitize import sanitize_docx_for_word

logger = logging.getLogger(__name__)

FIGURE_WIDTH_IN = 5.5


def render_blocks_document(
    doc_model: RenderDocument,
    context: dict,
    output_path: Path,
    *,
    style_shell_path: Path | None = None,
    skip_cover: bool = False,
    reference_layout_path: Path | None = None,
    content_start_marker: str = "Revision log",
    layout: str = "ep2737",
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    from ansys_report.report.ep2737_styles import open_document_from_shell

    doc = open_document_from_shell(style_shell_path)
    ep1581 = (layout or "").lower() == "ep1581"

    sections = doc_model.sections
    if skip_cover:
        sections = [section for section in sections if section.key != "cover"]

    for section in sections:
        _render_section(doc, section, context, ep1581=ep1581)

    if reference_layout_path is not None and reference_layout_path.exists():
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp_path = Path(tmp.name)
        try:
            doc.save(str(tmp_path))
            from ansys_report.report.ep2737_layout import merge_reference_front_matter

            merge_reference_front_matter(
                reference_layout_path,
                tmp_path,
                output_path,
                content_start_marker=content_start_marker,
            )
        finally:
            tmp_path.unlink(missing_ok=True)
    else:
        doc.save(str(output_path))

    sanitize_docx_for_word(output_path)
    if ep1581:
        from ansys_report.report.ep1581_docx_format import normalize_docx_portrait_a4

        normalize_docx_portrait_a4(output_path)
    logger.info("Wrote block-based report (%s): %s", layout, output_path)
    return output_path


def _render_section(doc: Document, section: RenderSection, context: dict, *, ep1581: bool) -> None:
    for block in section.blocks:
        _render_block(doc, block, context, ep1581=ep1581)


def _add_caption(doc: Document, text: str, *, ep1581: bool) -> None:
    if not text:
        return
    if ep1581:
        from ansys_report.report.ep1581_docx_format import add_caption as add_ep1581_caption

        add_ep1581_caption(doc, text)
        return
    cap = doc.add_paragraph(text)
    if cap.runs:
        cap.runs[0].italic = True


def _render_block(doc: Document, block: RenderBlock, context: dict, *, ep1581: bool) -> None:
    if block.kind == "cover":
        _render_cover(doc, context)
        doc.add_page_break()
        return

    if block.kind == "section_break" and ep1581:
        from ansys_report.report.ep1581_docx_format import add_section_break

        add_section_break(doc, block.orientation or "landscape")
        return

    if block.kind == "heading" and block.text:
        doc.add_heading(block.text, level=min(block.level, 4))
        return

    if block.kind == "paragraph" and block.text:
        if ep1581:
            from ansys_report.report.ep1581_docx_format import add_body_paragraph

            add_body_paragraph(doc, block.text)
        else:
            doc.add_paragraph(block.text)
        return

    if block.kind == "scalar" and block.label and block.value:
        text = f"{block.label}: {block.value}"
        if ep1581:
            from ansys_report.report.ep1581_docx_format import add_body_paragraph

            add_body_paragraph(doc, text)
        else:
            doc.add_paragraph(text)
        return

    if block.kind == "table":
        if block.caption:
            _add_caption(doc, block.caption, ep1581=ep1581)
        if ep1581:
            from ansys_report.report.ep1581_docx_format import add_table as add_ep1581_table

            add_ep1581_table(doc, block.headers, block.rows)
        else:
            _add_table(doc, block.headers, block.rows)
        return

    if block.kind == "figure" and block.image_path:
        try:
            if ep1581:
                from ansys_report.report.ep1581_docx_format import add_centered_picture

                add_centered_picture(doc, block.image_path, FIGURE_WIDTH_IN)
            else:
                doc.add_picture(block.image_path, width=Inches(FIGURE_WIDTH_IN))
        except Exception as exc:
            logger.warning("Could not embed figure %s: %s", block.slot, exc)
            if ep1581:
                raise RuntimeError(f"Figure embed failed: {block.caption or block.slot}") from exc
            doc.add_paragraph(f"[Figure unavailable: {block.caption or block.slot}]")
            return
        if block.caption:
            _add_caption(doc, block.caption, ep1581=ep1581)
        if block.description:
            if ep1581:
                from ansys_report.report.ep1581_docx_format import add_body_paragraph

                add_body_paragraph(doc, block.description)
            else:
                doc.add_paragraph(block.description)
        if ep1581:
            doc.add_paragraph("")
        return

    if block.kind == "pending":
        if ep1581 and not block.placeholder:
            raise RuntimeError(block.note or block.caption or "Pending content block")
        caption = block.caption or block.note or "Pending content"
        p = doc.add_paragraph()
        if caption.lower().startswith("table"):
            prefix = p.add_run("Table placeholder: ")
        else:
            prefix = p.add_run("Figure placeholder: ")
        prefix.bold = True
        body = p.add_run(caption)
        body.italic = True
        if block.note and block.note not in (caption, "Figure pending (skip_images)"):
            if ep1581:
                from ansys_report.report.ep1581_docx_format import add_body_paragraph

                add_body_paragraph(doc, block.note)
            else:
                doc.add_paragraph(block.note)
        return

    if block.kind == "narrative":
        for obs in block.observations:
            if ep1581:
                from ansys_report.report.ep1581_docx_format import add_body_paragraph

                add_body_paragraph(doc, str(obs))
            else:
                doc.add_paragraph(str(obs))
        for conclusion in block.conclusions:
            if ep1581:
                from ansys_report.report.ep1581_docx_format import add_body_paragraph

                add_body_paragraph(doc, str(conclusion))
            else:
                doc.add_paragraph(str(conclusion))
        for rec in block.recommendations:
            if ep1581:
                from ansys_report.report.ep1581_docx_format import add_body_paragraph

                add_body_paragraph(doc, str(rec))
            else:
                doc.add_paragraph(str(rec))
        return

    if block.kind == "page_break":
        doc.add_page_break()


def _render_cover(doc: Document, context: dict) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run(f"{context.get('bom_id', '')}\n")
    run.bold = True
    run.font.size = Pt(16)
    p.add_run(f"{context.get('title', '')}\n")
    doc.add_paragraph(f"Customer: {context.get('customer', '')}")
    prepared = context.get("prepared_by", {})
    checked = context.get("checked_by", {})
    approved = context.get("approved_by", {})
    doc.add_paragraph(
        f"Prepared by: {prepared.get('name', '')} ({prepared.get('role', '')}) | "
        f"Checked by: {checked.get('name', '')} | "
        f"Approved by: {approved.get('name', '')}"
    )


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    if not headers:
        return
    if not rows:
        doc.add_paragraph("(no data)")
        return
    col_count = max(len(headers), max((len(r) for r in rows), default=0))
    norm_headers = headers + [""] * (col_count - len(headers))
    table = doc.add_table(rows=1 + len(rows), cols=col_count)
    table.style = "Table Grid"
    for i, header in enumerate(norm_headers):
        table.rows[0].cells[i].text = header
    for r_idx, row in enumerate(rows, start=1):
        padded = row + [""] * (col_count - len(row))
        for c_idx, cell in enumerate(padded[:col_count]):
            table.rows[r_idx].cells[c_idx].text = cell
