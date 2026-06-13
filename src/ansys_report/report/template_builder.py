"""Create minimal EP1763-style docxtpl templates."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt


def create_ep2737_template(dest: Path) -> Path:
    """Phase 6 EP2737 report template: cover, equipment, modelling, modal, static, flange calcs."""
    doc = Document()
    _add_cover(doc)
    _add_ep2737_equipment(doc)
    _add_ep2737_modelling(doc)
    _add_modal_section(doc)
    _add_static_section(doc)
    _add_ep2737_flange_calcs(doc)
    doc.add_paragraph("Executive summary: {{ narrative.executive_summary }}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    from ansys_report.report.docx_sanitize import sanitize_docx_for_word

    sanitize_docx_for_word(dest)
    return dest


def _add_ep2737_equipment(doc: Document) -> None:
    doc.add_heading("Equipment Description", level=1)
    doc.add_paragraph("BOM ID: {{ equipment.bom_id }}")
    doc.add_paragraph("Title: {{ equipment.title }}")
    doc.add_paragraph(
        "{% if equipment.assembly %}Assembly mass: {{ equipment.assembly.mass_kg }} kg{% endif %}"
    )
    doc.add_paragraph(
        "{% for body in equipment.bodies %}"
        "{{ body.name }} | {{ body.material }} | {{ body.mass_kg }} kg\n"
        "{% endfor %}"
    )


def _add_ep2737_modelling(doc: Document) -> None:
    doc.add_heading("Modelling", level=1)
    doc.add_paragraph("Project: {{ modelling.project_name }}")
    doc.add_paragraph(
        "Mesh: {{ modelling.node_count }} nodes, {{ modelling.element_count }} elements"
    )
    doc.add_paragraph("Load steps: {{ modelling.load_steps }}")
    doc.add_paragraph("Contacts: {{ modelling.contacts | length }}")


def _add_ep2737_flange_calcs(doc: Document) -> None:
    doc.add_heading("Design Calculations - Flange Thickness (Table 25)", level=1)
    doc.add_paragraph(
        "{% for row in design_calcs.end_flange %}"
        "{{ row.label }} | {{ row.value }} {{ row.unit or '' }} | {{ row.verdict or '' }}\n"
        "{% endfor %}"
    )
    doc.add_heading("Bolt Preload (Table 27)", level=2)
    doc.add_paragraph(
        "{% for row in design_calcs.effort %}"
        "{{ row.label }} | {{ row.value }} {{ row.unit or '' }}\n"
        "{% endfor %}"
    )


def create_minimal_template_from_scratch(dest: Path) -> Path:
    """Build a minimal Jinja2 Word template when no sample docx is available."""
    doc = Document()
    _add_cover(doc)
    _add_section(doc, "Revision Log", "{{ sections.revision.body }}")
    _add_section(doc, "Scope", "{{ sections.scope.body }}")
    _add_section(doc, "Software and Tools", "ANSYS {{ ansys_version }}")
    _add_modal_section(doc)
    _add_static_section(doc)
    _add_design_calcs_section(doc)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    return dest


def create_minimal_template(source: Path, dest: Path) -> Path:
    """Copy sample docx and inject placeholder paragraphs (best-effort)."""
    try:
        from docxtpl import DocxTemplate

        tpl = DocxTemplate(str(source))
        tpl.save(str(dest))
        doc = Document(str(dest))
        doc.add_page_break()
        doc.add_heading("Automated Sections", level=1)
        doc.add_paragraph("BOM: {{ bom_id }} — {{ title }}")
        doc.save(str(dest))
        return dest
    except Exception:
        return create_minimal_template_from_scratch(dest)


def _add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run("{{ bom_id }}\n")
    run.bold = True
    run.font.size = Pt(16)
    p.add_run("{{ title }}\n")
    doc.add_paragraph("Customer: {{ customer }}")
    doc.add_paragraph(
        "Prepared by: {{ prepared_by.name }} ({{ prepared_by.role }}) | "
        "Checked by: {{ checked_by.name }} | Approved by: {{ approved_by.name }}"
    )
    doc.add_page_break()


def _add_section(doc: Document, heading: str, body_placeholder: str) -> None:
    doc.add_heading(heading, level=1)
    doc.add_paragraph(body_placeholder)


def _add_modal_section(doc: Document) -> None:
    doc.add_heading("Modal Analysis", level=1)
    doc.add_paragraph(
        "{% for mode in modal.modes %}"
        "Mode {{ mode.index }}: {{ mode.freq_hz }} Hz\n"
        "{% endfor %}"
    )
    doc.add_paragraph("{% for obs in modal.narrative.observations %}{{ obs }}\n{% endfor %}")
    doc.add_paragraph("{% for c in modal.narrative.conclusions %}{{ c }}\n{% endfor %}")


def _add_static_section(doc: Document) -> None:
    doc.add_heading("Static Structural", level=1)
    doc.add_paragraph("Max von-Mises stress: {{ static.max_stress_mpa }} MPa")
    doc.add_paragraph("Max deformation: {{ static.max_deformation_mm }} mm")
    doc.add_paragraph("FOS: {{ static.fos }}")
    doc.add_paragraph("{% for obs in static.narrative.observations %}{{ obs }}{% endfor %}")


def _add_design_calcs_section(doc: Document) -> None:
    doc.add_heading("Design Calculations", level=1)
    doc.add_paragraph("{% for row in design_calcs.bolt_load %}")
    doc.add_paragraph("{{ row.label }} | {{ row.value }} {{ row.unit }} | {{ row.verdict }}")
    doc.add_paragraph("{% endfor %}")
    doc.add_paragraph("Executive summary: {{ narrative.executive_summary }}")
